"""開発用の画面(/dev/)の API(バックログ B10。設計は docs/DEVTOOLS.md)。

利用先の Workbench で、CCP が読むデータ(BigQuery の 6 つの出力・ポリシーの文書)を用意するための道具。
BigQuery に問い合わせを投げられるのは、Workbench で動くこのサーバーだけなので、試しと答え合わせをここに置く。

  GET  /api/dev/status          有効か・ジョブのプロジェクト・アカウント・作業フォルダ
  GET  /api/dev/settings        作業の設定(よく見るデータセット・出す先 など。server/devwork/settings.json)
  PUT  /api/dev/settings
  GET  /api/dev/bq/datasets     ?project=  データセットの一覧
  GET  /api/dev/bq/tables       ?dataset=<プロジェクト>.<データセット>  表とビューの一覧
  GET  /api/dev/bq/table        ?ref=<プロジェクト>.<データセット>.<表>  列(入れ子は a.b の形に開く)・行数・パーティション
  GET  /api/dev/bq/preview      ?ref=…  先頭の行(list_rows。問い合わせではないので料金なし)
  POST /api/dev/bq/dryrun       {sql}  料金なしの試し。文の種類・読む量・結果の列
  POST /api/dev/bq/run          {sql, max_rows}  SELECT だけを、読む量の上限を付けて流す
  GET  /api/dev/bq/projects     見られるプロジェクト(データセットを、名前の一部で探すため)
  POST /api/dev/ai/chat         Gemini と会話して SQL を組み立てる。前提(仕様・手本・表の列・まとめ・いまの SQL)はサーバーで組み立てる。
                                思考・調べたこと・返事を、起きた順に NDJSON で返す
  GET  /api/dev/ai/models       使えるモデル(賢い順)と既定
  GET  /api/dev/work            6 つの出力の版の一覧・採用の版・出す先
  POST /api/dev/work/{出力}     SQL を新しい版として保存する   GET /api/dev/work/{出力}/{版}  版を開く
  PUT  /api/dev/work/{出力}/adopt  採用の版を決める          PUT /api/dev/work/target       出す先を決める
  POST /api/dev/check           答え合わせ(全件を流して bq_validate。run=false なら読む量の見積もりだけ)
  GET  /api/dev/gcs/list        ?bucket=&prefix=  GCS のフォルダ 1 段分(読むだけ)
  POST /api/dev/gcs/extract     選んだ文書を文字にする(Excel はシートごとに CSV の形。PDF は文字にしない)
  POST /api/dev/policy/convert  文書を POLICY_FORMAT の形に書き直す(NDJSON。PDF は Gemini にそのまま渡す)
  GET  /api/dev/policy/drafts   前に書き直したファイル(server/devwork/policies/)
  GET  /api/dev/files           作業ファイル(server/devwork/files/)  GET/PUT/DELETE /api/dev/files/{名前}  PUT …/pin  POST …/from-xlsx
  GET  /api/dev/chats           AI との会話の一覧   GET/PUT/DELETE /api/dev/chats/{id}  会話を開く・残す・消す(server/devwork/chats/)
  POST /api/dev/ai/context      Gemini に渡す前提を、そのまま返す(見る・コピーする)

決まり(DEVTOOLS.md の 6 章):
  - 環境変数 CCP_DEVTOOLS=1 のときだけ動く。それ以外は、画面(/dev/)も API も 404
  - BigQuery は読むだけ。dry-run で文の種類が SELECT のものだけを流す(DDL・DML・複数の文は流さない)
  - 読む量の上限は DEVTOOLS_MAX_GB(既定 10 GB)。超える問い合わせは BigQuery が止める
  - 作業ファイルは server/devwork/(Git に入れない。利用先の表の名前を含むため)。DB に置かないのは、import-pack で DB を置き換えても消えないように
"""
from __future__ import annotations

import asyncio
import base64
import json
import logging
import os
import re
import time
from datetime import date, datetime, time as dtime
from decimal import Decimal
from pathlib import Path
from typing import Any, AsyncIterator

from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import StreamingResponse
from pydantic import BaseModel, Field

from . import bq

log = logging.getLogger("ccp.devtools")

DEVWORK = Path(os.getenv("DEVTOOLS_DIR") or Path(__file__).resolve().parents[1] / "devwork")
# 画面に返す行の上限(多すぎるとブラウザが重くなる。全件の確かめは、別の段で行う)
MAX_ROWS = 1000
PREVIEW_ROWS = 20


def enabled() -> bool:
    return os.getenv("CCP_DEVTOOLS", "").strip().lower() in ("1", "true", "yes", "on")


def blocks(path: str) -> bool:
    """無効のとき、このパスを隠すか(画面の /dev/ と、API の /api/dev/)。main.py のミドルウェアから呼ぶ"""
    return not enabled() and (path == "/dev" or path.startswith("/dev/") or path.startswith("/api/dev/"))


def max_bytes() -> int:
    try:
        gb = float(os.getenv("DEVTOOLS_MAX_GB") or 10)
    except ValueError:
        gb = 10
    return int(gb * 1024**3)


async def require_enabled() -> None:
    if not enabled():
        raise HTTPException(404)


router = APIRouter(prefix="/api/dev", dependencies=[Depends(require_enabled)])


def plain(v: Any) -> Any:
    """BigQuery の値を、JSON にできる形にする(日時は ISO の文字、数値の Decimal は文字のまま桁を落とさない、バイト列は base64)"""
    if v is None or isinstance(v, (bool, int, float, str)):
        return v
    if isinstance(v, (datetime, date, dtime)):
        return v.isoformat()
    if isinstance(v, Decimal):
        return str(v)
    if isinstance(v, bytes):
        return base64.b64encode(v).decode()
    if isinstance(v, dict):
        return {k: plain(x) for k, x in v.items()}
    if isinstance(v, (list, tuple)):
        return [plain(x) for x in v]
    return str(v)


def flatten_fields(fields: Any, prefix: str = "") -> list[dict[str, Any]]:
    """表の列(SchemaField)を、入れ子(RECORD)を a.b の形に開いて並べる。Gemini に渡す列の一覧にもそのまま使う"""
    out: list[dict[str, Any]] = []
    for f in fields or []:
        name = f"{prefix}{f.name}"
        out.append({"name": name, "type": f.field_type, "mode": f.mode or "NULLABLE", "description": f.description or ""})
        if getattr(f, "fields", None):
            out.extend(flatten_fields(f.fields, f"{name}."))
    return out


def schema_columns(fields: Any) -> list[dict[str, str]]:
    return [{"name": f.name, "type": f.field_type} for f in (fields or [])]


def only_select(statement_type: str | None) -> str | None:
    """流してよい文か。だめなら、理由を返す(画面にそのまま出す)"""
    if statement_type == "SELECT":
        return None
    if not statement_type:
        return "文の種類が分かりませんでした(複数の文・スクリプトは流せません)"
    return f"SELECT ではないので流しません(文の種類: {statement_type})。表を作る SQL は、BigQuery のコンソールで実行してください"


class SqlIn(BaseModel):
    sql: str = Field(min_length=1, max_length=200_000)
    max_rows: int = Field(200, ge=1, le=MAX_ROWS)


def _error(e: Exception) -> dict[str, Any]:
    # BigQuery のエラーは、本文(message)に行と列の位置が入っている。Gemini に返して直させるので、そのまま渡す
    msg = getattr(e, "message", None) or str(e)
    errors = getattr(e, "errors", None) or []
    detail = "; ".join(x.get("message", "") for x in errors if isinstance(x, dict) and x.get("message")) if errors else ""
    return {"ok": False, "error": f"{type(e).__name__}: {detail or msg}"[:4000]}


def _dry_run(client: Any, sql: str) -> Any:
    from google.cloud import bigquery

    return client.query(sql, job_config=bigquery.QueryJobConfig(dry_run=True, use_query_cache=False))


# ---- 設定・状態

def _settings_path() -> Path:
    return DEVWORK / "settings.json"


@router.get("/status")
async def status() -> dict[str, Any]:
    try:
        account = await asyncio.to_thread(bq.account)
    except Exception as e:  # noqa: BLE001 — 認証がない環境(手元)でも、画面は開けるように
        account = f"(認証を読めません: {type(e).__name__})"
    return {"enabled": True, "project": bq.job_project(), "account": account, "devwork": str(DEVWORK), "max_gb": max_bytes() / 1024**3}


@router.get("/settings")
async def get_settings() -> dict[str, Any]:
    p = _settings_path()
    if not p.is_file():
        return {}
    try:
        return json.loads(p.read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return {}


@router.put("/settings")
async def put_settings(body: dict[str, Any]) -> dict[str, Any]:
    DEVWORK.mkdir(parents=True, exist_ok=True)
    _settings_path().write_text(json.dumps(body, ensure_ascii=False, indent=2), encoding="utf-8")
    return {"ok": True}


# ---- BigQuery を見る(メタデータ・先頭の行。どれも料金なし)

@router.get("/bq/datasets")
async def datasets(project: str | None = None) -> dict[str, Any]:
    def go() -> dict[str, Any]:
        c = bq.client()
        p = project or c.project
        return {"ok": True, "project": p, "datasets": sorted(d.dataset_id for d in c.list_datasets(project=p))}

    try:
        return await asyncio.to_thread(go)
    except Exception as e:  # noqa: BLE001
        return _error(e)


@router.get("/bq/tables")
async def tables(dataset: str) -> dict[str, Any]:
    def go() -> dict[str, Any]:
        c = bq.client()
        rows = []
        for t in c.list_tables(dataset.strip().strip("`")):
            rows.append({"id": t.table_id, "ref": f"{t.project}.{t.dataset_id}.{t.table_id}", "type": t.table_type})
        return {"ok": True, "tables": sorted(rows, key=lambda r: r["id"])}

    try:
        return await asyncio.to_thread(go)
    except Exception as e:  # noqa: BLE001
        return _error(e)


def _sources_or_empty(t: Any) -> list[str]:
    """ビューなら参照している元の表(画面の「参照している元の表」)。分からなくても、表の説明は返す"""
    if not t.view_query:
        return []
    try:
        return view_sources(bq.client(), f"{t.project}.{t.dataset_id}.{t.table_id}")
    except Exception:  # noqa: BLE001
        return []


@router.get("/bq/table")
async def table(ref: str) -> dict[str, Any]:
    def go() -> dict[str, Any]:
        t = bq.client().get_table(ref.strip().strip("`"))
        part = t.time_partitioning
        return {
            "ok": True,
            "ref": f"{t.project}.{t.dataset_id}.{t.table_id}",
            "type": t.table_type,
            "description": t.description or "",
            "rows": t.num_rows,
            "bytes": t.num_bytes,
            "modified": plain(t.modified),
            "location": t.location,
            "partitioning": ({"type": part.type_, "field": part.field or "_PARTITIONTIME"} if part else None),
            "clustering": list(t.clustering_fields or []),
            "view_sql": t.view_query or "",
            "sources": _sources_or_empty(t),
            "columns": flatten_fields(t.schema),
        }

    try:
        return await asyncio.to_thread(go)
    except Exception as e:  # noqa: BLE001
        return _error(e)


@router.get("/bq/preview")
async def preview(ref: str) -> dict[str, Any]:
    def go() -> dict[str, Any]:
        c = bq.client()
        t = c.get_table(ref.strip().strip("`"))
        if t.table_type != "TABLE":
            # ビューは list_rows で読めない。先頭の行は、SQL の試しで見る
            return {"ok": False, "error": f"{t.table_type} は先頭の行を読めません(SQL の試しで SELECT * … LIMIT 20 を流してください)"}
        it = c.list_rows(t, max_results=PREVIEW_ROWS)
        rows = [plain(dict(r.items())) for r in it]
        return {"ok": True, "columns": schema_columns(t.schema), "rows": rows}

    try:
        return await asyncio.to_thread(go)
    except Exception as e:  # noqa: BLE001
        return _error(e)


# ---- SQL を試す

def price_per_tib() -> float:
    """オンデマンドの料金(1 TiB あたりのドル)。場所や契約で違うので、設定(DEVTOOLS_PRICE_PER_TIB)で変えられる"""
    try:
        return float(os.getenv("DEVTOOLS_PRICE_PER_TIB") or 6.25)
    except ValueError:
        return 6.25


def cost_usd(n: int | None) -> float | None:
    return None if n is None else round(n / 1024**4 * price_per_tib(), 4)


def weight(processed: int | None, slot_ms: int | None) -> str:
    """重さの目安(読む量とスロット時間から。画面に一言で出す)"""
    gb = (processed or 0) / 1024**3
    slot_s = (slot_ms or 0) / 1000
    if gb >= 100 or slot_s >= 3600:
        return "重い"
    if gb >= 1 or slot_s >= 60:
        return "ふつう"
    return "軽い"


# SQL の中の `プロジェクト.データセット.表` の形(参照したビューを見つけるため)
TABLE_REF = re.compile(r"`?([A-Za-z0-9_-]+\.[A-Za-z0-9_]+\.[A-Za-z0-9_$*]+)`?")


def referenced(c: Any, sql: str, job: Any) -> dict[str, list[str]]:
    """参照した表とビュー。BigQuery のジョブの記録には、ビューの中身(元の表)しか残らないので、
    SQL に書いた名前のうち、元の表に出てこないものを調べて、ビューを見つける"""
    tables = [f"{t.project}.{t.dataset_id}.{t.table_id}" for t in (job.referenced_tables or [])]
    views: list[str] = []
    for name in dict.fromkeys(TABLE_REF.findall(sql)):
        if name in tables or "*" in name:
            continue
        try:
            t = c.get_table(name)
        except Exception:  # noqa: BLE001 — 列の名前などが、たまたま同じ形だっただけ
            continue
        if t.table_type in ("VIEW", "MATERIALIZED_VIEW"):
            views.append(f"{t.project}.{t.dataset_id}.{t.table_id}")
    return {"tables": list(dict.fromkeys(tables)), "views": views}


def job_info(c: Any, sql: str, job: Any) -> dict[str, Any]:
    """ジョブの実行情報(重いのか・高いのか・どこを読んだのか)。画面の結果の下に出す"""
    elapsed = (job.ended - job.started).total_seconds() if job.started and job.ended else None
    try:
        refs = referenced(c, sql, job)
    except Exception:  # noqa: BLE001
        refs = {"tables": [], "views": []}
    return {
        "job_id": job.job_id,
        "location": job.location,
        "processed": job.total_bytes_processed,
        "billed": job.total_bytes_billed,
        "cost_usd": cost_usd(job.total_bytes_billed),
        "price_per_tib": price_per_tib(),
        "slot_ms": job.slot_millis,
        "elapsed": elapsed,
        "cached": bool(job.cache_hit),
        "weight": weight(job.total_bytes_processed, job.slot_millis),
        **refs,
    }


@router.post("/bq/dryrun")
async def dryrun(body: SqlIn) -> dict[str, Any]:
    def go() -> dict[str, Any]:
        c = bq.client()
        job = _dry_run(c, body.sql)
        st = job.statement_type
        try:
            refs = referenced(c, body.sql, job)
        except Exception:  # noqa: BLE001
            refs = {"tables": [], "views": []}
        return {"ok": True, "statement_type": st, "bytes": job.total_bytes_processed, "columns": schema_columns(job.schema),
                "runnable": only_select(st) is None, "reason": only_select(st), "max_bytes": max_bytes(),
                # 流したときの概算(課金は最小 10 MB から。キャッシュに当たれば 0)
                "cost_usd": cost_usd(max(job.total_bytes_processed or 0, 10 * 1024**2)), "price_per_tib": price_per_tib(),
                "weight": weight(job.total_bytes_processed, None), **refs}

    try:
        return await asyncio.to_thread(go)
    except Exception as e:  # noqa: BLE001
        return _error(e)


@router.post("/bq/run")
async def run(body: SqlIn) -> dict[str, Any]:
    def go() -> dict[str, Any]:
        from google.cloud import bigquery

        c = bq.client()
        dry = _dry_run(c, body.sql)
        if (why := only_select(dry.statement_type)) is not None:
            return {"ok": False, "error": why}
        t0 = time.monotonic()
        job = c.query(body.sql, job_config=bigquery.QueryJobConfig(maximum_bytes_billed=max_bytes()))
        it = job.result(max_results=body.max_rows)
        rows = [plain(dict(r.items())) for r in it]
        return {
            "ok": True,
            "columns": schema_columns(it.schema),
            "rows": rows,
            "total_rows": it.total_rows,
            "bytes": job.total_bytes_processed,
            "billed": job.total_bytes_billed,
            "seconds": round(time.monotonic() - t0, 2),
            "cached": bool(job.cache_hit),
            "job": job_info(c, body.sql, job),
        }

    try:
        return await asyncio.to_thread(go)
    except Exception as e:  # noqa: BLE001
        return _error(e)


@router.get("/bq/projects")
async def projects() -> dict[str, Any]:
    """このアカウントが見られるプロジェクト(データセットを、プロジェクトの名前の一部でも探せるように)"""
    def go() -> dict[str, Any]:
        c = bq.client()
        ids = sorted({p.project_id for p in c.list_projects(max_results=500)} | ({c.project} if c.project else set()))
        return {"ok": True, "default": c.project, "projects": ids}

    try:
        return await asyncio.to_thread(go)
    except Exception as e:  # noqa: BLE001
        return _error(e)


# ---- Gemini と SQL を組み立てる(DEVTOOLS.md の 4 章)
# 前提(system_instruction)は、毎回サーバーで組み立て直す。画面から送るのは、会話と「何を渡すか」の選択だけ。
# そうすると、表の列が変わっても、仕様の文書を直しても、次の発言から新しい前提で答える。
# Gemini には、BigQuery を調べる道具(関数)を渡す。「A データセットの ctlg を含む表から…」のように頼めば、
# Gemini が自分で表を探し、列を見て、SQL を確かめてから答える(利用者が表を 1 つずつ選んで渡さなくてよいように)。

DOCS = Path(__file__).resolve().parents[2] / "docs"
OUTPUT_NAMES = ("spu", "sku", "sku_stock", "weekly", "forecast", "ranking")
# 先頭の行を Gemini に渡すとき、外す列(個人を特定できる・自由な文章の列。名前で見分ける)
PRIVATE = re.compile(r"(name|_nm$|mail|phone|tel|addr|birth|zip|postal|review|comment|memo|note|text|txt|氏名|名前|住所|電話|メール|誕生|郵便|レビュー|コメント)", re.I)
SAMPLE_ROWS = 5
# Gemini が道具を使う回数の上限(1 回の発言につき)。表を探す → 列を見る → 確かめる、で足りる数
MAX_TOOL_ROUNDS = 16
# Gemini が流す SELECT の、返す行の上限
TOOL_ROWS = 20


def chapters(md: str, numbers: tuple[int, ...]) -> str:
    """Markdown から「## <番号>. …」の章だけを抜き出す"""
    out: list[str] = []
    for part in re.split(r"(?m)^(?=## )", md):
        m = re.match(r"## (\d+)\.", part)
        if m and int(m.group(1)) in numbers:
            out.append(part.strip().rstrip("-").strip())
    return "\n\n".join(out)


def _doc(name: str) -> str:
    try:
        return (DOCS / name).read_text(encoding="utf-8")
    except OSError:
        return ""


ROLE = """あなたは BigQuery(GoogleSQL)の SQL を書く専門家です。利用者と会話しながら、SQL を組み立てます。
利用者は、Campaign Contents Planner(CCP)が読むデータ(下の「CCP が読む出力の仕様」の 6 つの出力)を、利用先の BigQuery の元の表から作ろうとしています。
元の表の列を、CCP のどの列に当てればよいかを一緒に考えるのも、あなたの仕事です。

BigQuery を調べる道具(関数)を使えます:
- list_datasets: プロジェクトのデータセットの一覧
- list_tables: データセットの表の一覧(name_contains で名前の一部で絞れる。大文字・小文字は区別しない)
- describe_table: 表の列・型・説明・パーティションと、先頭の数行
- dry_run: SQL を料金なしで確かめる(文法・列の有無・読む量)
- describe_view_sources: ビューが参照している元の表(ビューの中のビューもたどる)と、それぞれの列の定義
- run_select: SELECT を流して、先頭の行を見る(値の種類を確かめる SELECT DISTINCT など。読む量に上限がある)
推測で書かず、道具で表と列を確かめてから SQL を書いてください。SQL を返す前に、dry_run で通ることを確かめてください。
利用者が名前の一部しか言わないとき(「ctlg が含まれた表」など)は、list_tables で探してください。
ビューは、ビューの SQL(describe_table で見える)と、参照している元の表の列(describe_view_sources)をたどり、
CCP の列に合う元の列を探してください。ビューに出ていない列でも、元の表にあれば候補にし、どのビューに足せばよいか・元の表から直接読むべきかを考えてください。

最後の返答は、次の形の JSON だけにしてください(```json で囲まない):
{"reply": "利用者への返事(日本語。Markdown の見出し・箇条書き・**太字**・表を使ってよい。SQL はここに書かない)",
 "sqls": [{"title": "何の SQL か(短く)", "sql": "SQL の全文"}],
 "questions": ["決めきれなかったこと・利用者に確かめたいこと(なければ空の配列)"]}
sqls は、新しく書いた・直した SQL があるときだけ(なければ空の配列)。直すときも、部分ではなく、そのまま流せる全文。
値を確かめる SQL と、本番の SQL のように、いくつか出すときは並べる(利用者は 1 つずつ試す)。

決まり:
- 方言は GoogleSQL(BigQuery)。表は `プロジェクト.データセット.表` の形で、バッククォートで囲む
- 確かめた表と列だけを使う。ない列を推測で書かない(足りなければ、questions で聞く)
- CCP の列に当てるときは、仕様のとおりに AS で付け替える。当て方を表にして reply に書く(元の列 → CCP の列、変換、自信のなさ)
- 読む量を抑える(パーティションの列があれば、期間で絞る。SELECT * を避ける)
- 「いまの SQL」とエラーが渡されたら、まずエラーの原因を reply で短く説明し、直した全文を sql に入れる
- 「いつも渡すファイル」にコーディングルールがあれば、SQL はそれに必ず従う(名前の付け方・書き方・コメント など)。従えないところは reply で理由を書く
- 「いつも渡すファイル」に引き継ぎがあれば、その続きとして話す(決まったことを、もう一度聞かない)

作業ファイル(Workbench のフォルダ。利用者も画面で見られる)を使う道具もあります:
- list_files・read_file: 作業ファイルの一覧と中身
- save_file: 作業ファイルに保存する(.md / .sql / .txt / .csv)。利用者に「保存して」と頼まれたとき、引き継ぎを作るときに使う。
  同じ名前があると上書きするので、上書きしてよいか分からないときは、別の名前にする
"""

TABLE_POLICY = """## 出力の方針(利用者が決めたこと)
- 出力は、ビューではなく表にする。最終的に `CREATE OR REPLACE TABLE <出す先> AS SELECT …` で作り、BigQuery のスケジュールされたクエリで定期的に更新する
- ただし、この会話で書くのは **SELECT の部分だけ**(この画面では SELECT しか流せない。表を作る形は、合格したあとで作る)
- 何度流しても同じ結果になる書き方にする
- `as_of` には元のデータの時点を入れる(CURRENT_TIMESTAMP() にしない)
- 出す先は、元の表と同じ場所(リージョン)
"""


class ChatTurn(BaseModel):
    role: str = Field(pattern="^(user|model)$")
    text: str = Field(max_length=200_000)


class ChatIn(BaseModel):
    messages: list[ChatTurn] = Field(default_factory=list, max_length=200)
    # 使うモデル(画面で選ぶ)。なければ既定(default_model)
    model: str | None = Field(None, max_length=200)
    # 作ろうとしている出力(spu など)。手本の SQL を渡す。なしでも、仕様はいつも全文を渡す
    output: str | None = None
    # 画面で開いているデータセット(「このデータセット」と言ったときの手がかり)
    dataset: str = ""
    tables: list[str] = Field(default_factory=list, max_length=20)
    sample_rows: bool = True
    # Gemini に SELECT を流させるか(読む量の上限つき)。切ると dry_run だけ
    allow_run: bool = True
    notes: str = Field("", max_length=200_000)
    sql: str = Field("", max_length=200_000)
    last: str = Field("", max_length=50_000)


def _sample(c: Any, t: Any) -> list[str]:
    """先頭の数行(Markdown の表)。名前から個人の情報・自由な文章と見られる列は外す"""
    if t.table_type != "TABLE":
        return []
    try:
        rows = [plain(dict(r.items())) for r in c.list_rows(t, max_results=SAMPLE_ROWS)]
    except Exception:  # noqa: BLE001 — 行が読めなくても、列だけで続ける
        return []
    cols = [f.name for f in t.schema if not PRIVATE.search(f.name)]
    hidden = [f.name for f in t.schema if PRIVATE.search(f.name)]
    if not rows or not cols:
        return []
    out = ["", f"先頭の {len(rows)} 行(値の形を見るため。外した列: {', '.join(hidden) or 'なし'}):", "", "| " + " | ".join(cols) + " |", "|" + "---|" * len(cols)]
    out += ["| " + " | ".join(str(r.get(k, ""))[:60].replace("|", " ").replace("\n", " ") for k in cols) + " |" for r in rows]
    return out


def view_sources(c: Any, ref: str) -> list[str]:
    """ビューが参照している元の表(料金なしの dry-run で BigQuery に聞く。ビューの中のビューも、元の表までたどった結果が返る)"""
    job = _dry_run(c, f"SELECT * FROM `{ref}`")
    me = ref.replace(":", ".")
    out = []
    for t in job.referenced_tables or []:
        name = f"{t.project}.{t.dataset_id}.{t.table_id}"
        if name != me and name not in out:
            out.append(name)
    return out


def describe(c: Any, ref: str, with_rows: bool) -> str:
    """表 1 つの説明(Gemini に渡す形)"""
    try:
        t = c.get_table(ref.strip().strip("`"))
    except Exception as e:  # noqa: BLE001 — 読めない表は、読めないことを伝えて続ける
        return f"### `{ref}`\n(読めませんでした: {type(e).__name__}: {e})"[:1000]
    part = t.time_partitioning
    lines = [f"### `{t.project}.{t.dataset_id}.{t.table_id}`({t.table_type}、約 {t.num_rows or 0:,} 行、{t.location})"]
    if t.description:
        lines.append(t.description)
    if part:
        lines.append(f"パーティション: {part.field or '_PARTITIONTIME'}({part.type_})")
    if t.clustering_fields:
        lines.append(f"クラスタ: {', '.join(t.clustering_fields)}")
    if t.view_query:
        lines += ["", "ビューの SQL(DDL の本体):", "```sql", t.view_query.strip()[:8000], "```"]
        try:
            src = view_sources(c, f"{t.project}.{t.dataset_id}.{t.table_id}")
            if src:
                lines += ["", "このビューが参照している元の表(列は describe_view_sources で見られる): " + ", ".join(f"`{x}`" for x in src)]
        except Exception:  # noqa: BLE001 — 参照先が分からなくても、ビューの SQL は渡せる
            pass
    lines += ["", "| 列 | 型 | モード | 説明 |", "|---|---|---|---|"]
    lines += [f"| {f['name']} | {f['type']} | {f['mode']} | {f['description'].replace('|', ' ')} |" for f in flatten_fields(t.schema)]
    if with_rows:
        lines += _sample(c, t)
    return "\n".join(lines)


def build_context(body: ChatIn, client: Any | None) -> str:
    """Gemini に渡す前提(DEVTOOLS.md の 4 章の順。上ほど変わらない内容)"""
    parts = [ROLE, TABLE_POLICY]
    pinned = pinned_context()
    if pinned:
        parts.append(pinned)
    if client is not None:
        where = [f"- 接続しているプロジェクト(問い合わせを走らせるプロジェクト): `{client.project}`"]
        if body.dataset:
            where.append(f"- 画面で開いているデータセット: `{body.dataset}`")
        parts.append("## いまの接続\n\n" + "\n".join(where))
    parts.append("## CCP が読む出力の仕様(BIGQUERY_EXPORT.md)\n\n" + chapters(_doc("BIGQUERY_EXPORT.md"), tuple(range(1, 20))))
    rules = chapters(_doc("BIGQUERY_CONVERT.md"), (4,))
    if rules:
        parts.append(rules)
    if body.output in OUTPUT_NAMES:
        example = DOCS / "bq_example" / f"{body.output}.sql"
        parts.append(f"## いま作ろうとしている出力: {body.output}")
        if example.is_file():
            parts.append("## 手本の SQL(練習用の元データ `{src}` から作った例。元の表の名前・列は、利用先と違う)\n\n```sql\n" + example.read_text(encoding="utf-8") + "\n```")
    if body.tables and client is not None:
        parts.append("## 利用者が選んで渡した表\n\n" + "\n\n".join(describe(client, r, body.sample_rows) for r in body.tables))
    if body.notes.strip():
        parts.append("## 仕様のまとめ(利用先の資料から、利用者がまとめたもの)\n\n" + body.notes.strip())
    if body.sql.strip():
        parts.append("## いまの SQL(画面のエディタの中身)\n\n```sql\n" + body.sql.strip() + "\n```")
    if body.last.strip():
        parts.append("## いまの SQL を流した結果(エラー、または結果の先頭)\n\n" + body.last.strip())
    return "\n\n".join(parts)


def parse_reply(text: str) -> dict[str, Any]:
    """Gemini の最後の返答(JSON)を読む。```json で囲まれていても、前後に文があっても拾う。崩れていたら、本文をそのまま返事にする"""
    candidates = [text]
    m = re.search(r"```(?:json)?\s*(\{.*\})\s*```", text, re.S)
    if m:
        candidates.append(m.group(1))
    a, b = text.find("{"), text.rfind("}")
    if 0 <= a < b:
        candidates.append(text[a : b + 1])
    for c in candidates:
        try:
            v = json.loads(c)
        except ValueError:
            continue
        if isinstance(v, dict) and ("reply" in v or "sql" in v or "sqls" in v):
            q = v.get("questions")
            sqls = [{"title": str(x.get("title") or ""), "sql": str(x.get("sql") or "").strip()} for x in (v.get("sqls") or []) if isinstance(x, dict)]
            # 1 つだけの形(sql)で返してきたときも受ける
            if isinstance(v.get("sql"), str) and v["sql"].strip():
                sqls.insert(0, {"title": "", "sql": v["sql"].strip()})
            return {"reply": str(v.get("reply") or ""), "sqls": [x for x in sqls if x["sql"]],
                    "questions": [str(x) for x in q if x] if isinstance(q, list) else []}
    return {"reply": text, "sqls": [], "questions": []}


# ---- Gemini が使う道具(どれも読むだけ。BigQuery に書き込むものはない)

def tool_declarations() -> list[Any]:
    from google.genai import types

    S = types.Schema
    T = types.Type
    return [types.Tool(function_declarations=[
        types.FunctionDeclaration(name="list_datasets", description="プロジェクトのデータセットの一覧。project を省くと、接続しているプロジェクト",
                                  parameters=S(type=T.OBJECT, properties={"project": S(type=T.STRING)})),
        types.FunctionDeclaration(name="list_tables", description="データセットの表・ビューの一覧。name_contains で、名前の一部(大文字・小文字は区別しない)で絞れる",
                                  parameters=S(type=T.OBJECT, properties={"dataset": S(type=T.STRING, description="プロジェクト.データセット、またはデータセットだけ(接続しているプロジェクト)"),
                                                                          "name_contains": S(type=T.STRING)}, required=["dataset"])),
        types.FunctionDeclaration(name="describe_table", description="表の列・型・説明・パーティション・行数と、先頭の数行",
                                  parameters=S(type=T.OBJECT, properties={"table": S(type=T.STRING, description="プロジェクト.データセット.表")}, required=["table"])),
        types.FunctionDeclaration(name="describe_view_sources", description="ビューが参照している元の表(ビューの中のビューもたどる)と、それぞれの列の定義",
                                  parameters=S(type=T.OBJECT, properties={"view": S(type=T.STRING, description="プロジェクト.データセット.ビュー")}, required=["view"])),
        types.FunctionDeclaration(name="dry_run", description="SQL を料金なしで確かめる。文の種類・読む量・結果の列、またはエラー",
                                  parameters=S(type=T.OBJECT, properties={"sql": S(type=T.STRING)}, required=["sql"])),
        types.FunctionDeclaration(name="list_files", description="作業ファイルの一覧(名前・大きさ・更新日時・いつも渡すか)",
                                  parameters=S(type=T.OBJECT, properties={})),
        types.FunctionDeclaration(name="read_file", description="作業ファイルの中身",
                                  parameters=S(type=T.OBJECT, properties={"name": S(type=T.STRING)}, required=["name"])),
        types.FunctionDeclaration(name="save_file", description="作業ファイルに保存する(.md / .sql / .txt / .csv)。同じ名前は上書き",
                                  parameters=S(type=T.OBJECT, properties={"name": S(type=T.STRING, description="ファイル名(例: 引き継ぎ_20261023.md、spu_from_ctlg.sql)"),
                                                                          "content": S(type=T.STRING)}, required=["name", "content"])),
        types.FunctionDeclaration(name="run_select", description=f"SELECT を流して、先頭 {TOOL_ROWS} 行を見る(値の種類・件数の確かめ)。SELECT 以外と、読む量が上限を超えるものは流れない",
                                  parameters=S(type=T.OBJECT, properties={"sql": S(type=T.STRING)}, required=["sql"])),
    ])]


def _qualify(c: Any, name: str) -> str:
    """データセットだけなら、接続しているプロジェクトを付ける"""
    name = name.strip().strip("`")
    return name if "." in name else f"{c.project}.{name}"


def run_tool(c: Any, name: str, args: dict[str, Any], body: ChatIn) -> tuple[dict[str, Any], str]:
    """道具を 1 つ使う。Gemini に返す結果と、画面に出す 1 行(何をしたか)を返す"""
    from google.cloud import bigquery

    try:
        if name == "list_datasets":
            p = (args.get("project") or c.project).strip()
            ids = sorted(d.dataset_id for d in c.list_datasets(project=p))
            return {"project": p, "datasets": ids}, f"データセットの一覧を見た({p}、{len(ids)} 件)"
        if name == "list_tables":
            ds = _qualify(c, str(args.get("dataset") or body.dataset))
            key = str(args.get("name_contains") or "").lower()
            rows = [{"table": f"{t.project}.{t.dataset_id}.{t.table_id}", "type": t.table_type} for t in c.list_tables(ds) if key in t.table_id.lower()]
            more = len(rows) > 300
            return {"dataset": ds, "tables": rows[:300], "truncated": more}, f"表を探した({ds}{f'、「{key}」を含む' if key else ''}、{len(rows)} 件)"
        if name == "describe_table":
            ref = _qualify(c, str(args.get("table") or ""))
            return {"description": describe(c, ref, body.sample_rows)}, f"表の列を見た({ref})"
        if name == "describe_view_sources":
            ref = _qualify(c, str(args.get("view") or ""))
            src = view_sources(c, ref)[:12]
            # 元の表は列だけ(先頭の行は、必要なら describe_table で個別に見る。量を抑えるため)
            return {"view": ref, "sources": [describe(c, x, False) for x in src]}, f"ビューの参照先を見た({ref}、元の表 {len(src)} 件)"
        if name == "dry_run":
            job = c.query(str(args.get("sql") or ""), job_config=bigquery.QueryJobConfig(dry_run=True, use_query_cache=False))
            return ({"statement_type": job.statement_type, "bytes": job.total_bytes_processed, "columns": schema_columns(job.schema)},
                    f"SQL を確かめた(通った。読む量 {job.total_bytes_processed or 0:,} バイト)")
        if name == "list_files":
            return {"files": list_work_files()}, "作業ファイルの一覧を見た"
        if name == "read_file":
            n = str(args.get("name") or "")
            return {"name": n, "content": read_work_file(n)}, f"作業ファイルを読んだ({n})"
        if name == "save_file":
            n = str(args.get("name") or "")
            existed = write_work_file(n, str(args.get("content") or ""))
            return {"saved": n, "overwritten": existed}, f"作業ファイルに保存した({n}{'。上書き' if existed else ''})"
        if name == "run_select":
            if not body.allow_run:
                return {"error": "利用者が、SQL を流すことを止めています。dry_run で確かめてください"}, "SELECT を流そうとした(止めてある)"
            sql = str(args.get("sql") or "")
            dry = _dry_run(c, sql)
            if (why := only_select(dry.statement_type)) is not None:
                return {"error": why}, "SELECT 以外を流そうとした(止めた)"
            job = c.query(sql, job_config=bigquery.QueryJobConfig(maximum_bytes_billed=max_bytes()))
            it = job.result(max_results=TOOL_ROWS)
            rows = [plain(dict(r.items())) for r in it]
            return ({"columns": schema_columns(it.schema), "rows": rows, "total_rows": it.total_rows},
                    f"SELECT を流した({it.total_rows or 0:,} 行、読んだ量 {job.total_bytes_processed or 0:,} バイト)")
        return {"error": f"知らない道具です: {name}"}, f"知らない道具 {name}"
    except Exception as e:  # noqa: BLE001 — エラーも Gemini に返して、直させる
        err = _error(e)["error"]
        return {"error": err}, f"{name} がエラーになった({err[:120]})"


async def _context(body: ChatIn) -> tuple[str, Any]:
    client = await asyncio.to_thread(bq.client)
    return await asyncio.to_thread(build_context, body, client), client


def _version(name: str) -> tuple[float, int]:
    """モデルの名前から、新しさを比べる値(gemini-3.5-pro → 3.5。preview は、同じ版の正式より後ろ)"""
    m = re.search(r"gemini-(\d+(?:\.\d+)?)", name)
    return (float(m.group(1)) if m else 0.0, 0 if "preview" in name else 1)


def pick_models(names: list[str]) -> list[str]:
    """Vertex AI のモデルの一覧から、会話に使える Gemini を、賢い順(Pro → Flash、新しい版が先)に並べる"""
    keep = []
    for n in names:
        n = n.split("/")[-1]
        if not n.startswith("gemini-") or not re.search(r"-(pro|flash)\b", n):
            continue
        if re.search(r"(image|tts|audio|live|embedding|lite|exp|thinking-exp)", n):
            continue
        keep.append(n)
    return sorted(set(keep), key=lambda n: (0 if "-pro" in n else 1, tuple(-x for x in _version(n))))


_models_cache: list[str] | None = None


async def available_models() -> list[str]:
    """使えるモデル(Vertex AI に聞く。聞けなければ設定のモデルだけ)。プロセスの間は覚えておく"""
    global _models_cache
    if _models_cache is None:
        from . import ai

        names: list[str] = []
        try:
            pager = await ai._get_client().aio.models.list(config={"query_base": True, "page_size": 200})
            async for m in pager:
                names.append(m.name or "")
        except Exception:  # noqa: BLE001 — 一覧を出せない環境でも、設定のモデルで動く
            log.info("Vertex AI のモデルの一覧を読めませんでした", exc_info=True)
        _models_cache = pick_models(names)
    pinned = [m for m in (os.getenv("DEVTOOLS_GEMINI_MODEL"), _pro_guess or None, os.getenv("GEMINI_MODEL")) if m]
    return list(dict.fromkeys(pinned[:-1] + _models_cache + pinned[-1:]))


async def default_model() -> str:
    """開発用の画面の既定のモデル。設定(DEVTOOLS_GEMINI_MODEL)があればそれ、なければ一覧でいちばん賢いもの。
    SQL を組み立てるのは、配信の下書きと違って 1 回ずつの呼び出しなので、速さより賢さを取る"""
    global _pro_guess
    if os.getenv("DEVTOOLS_GEMINI_MODEL"):
        return os.environ["DEVTOOLS_GEMINI_MODEL"]
    models = await available_models()
    pros = [m for m in (_models_cache or []) if "-pro" in m]
    if pros:
        return pros[0]
    # 一覧を読めない環境: CCP のモデル(GEMINI_MODEL)の Flash を Pro に置き換えた名前が、使えるかを 1 度だけ聞く
    base = os.getenv("GEMINI_MODEL", "gemini-3.5-flash")
    if _pro_guess is None:
        _pro_guess = ""
        guess = pro_of(base)
        if guess:
            from . import ai

            try:
                await ai._get_client().aio.models.get(model=guess)
                _pro_guess = guess
            except Exception:  # noqa: BLE001 — その Pro がない・権限がない。CCP のモデルのまま
                log.info("%s は使えませんでした。%s を使います", guess, base)
    if _pro_guess:
        return _pro_guess
    return models[0] if models else base


def pro_of(name: str) -> str | None:
    """Flash の名前から、同じ版の Pro の名前を作る(gemini-3.8-flash → gemini-3.8-pro)。作れなければ None"""
    guess = re.sub(r"-flash(-lite)?", "-pro", name, count=1)
    return guess if guess != name else None


_pro_guess: str | None = None


@router.get("/ai/models")
async def models() -> dict[str, Any]:
    return {"ok": True, "default": await default_model(), "models": await available_models(),
            "thinking": (os.getenv("DEVTOOLS_THINKING") or "high").lower()}


def _line(obj: dict[str, Any]) -> bytes:
    return (json.dumps(obj, ensure_ascii=False) + "\n").encode("utf-8")


async def chat_events(body: ChatIn) -> AsyncIterator[dict[str, Any]]:
    """会話の 1 回分を、起きた順に返す(画面が、思考・調べたこと・返事を、届いたそばから出せるように)。
      {"type": "start", "model": …}       {"type": "thought", "text": …}(思考の要約。少しずつ届く)
      {"type": "step", "text": …}         {"type": "final", "reply": …, "sqls": […], "questions": […]}
      {"type": "error", "error": …}"""
    from google.genai import types

    from . import ai

    if not body.messages or body.messages[-1].role != "user":
        yield {"type": "error", "error": "送る発言がありません"}
        return
    model = body.model or await default_model()
    level = (os.getenv("DEVTOOLS_THINKING") or "high").upper()
    yield {"type": "start", "model": model, "thinking": level.lower()}
    try:
        context, client = await _context(body)
    except Exception as e:  # noqa: BLE001
        yield {"type": "error", "error": _error(e)["error"]}
        return
    config = types.GenerateContentConfig(
        system_instruction=context,
        tools=tool_declarations(),
        # 道具の呼び出しは、ここで 1 つずつ行う(画面に「何をしたか」を出し、回数に上限を付けるため)
        automatic_function_calling=types.AutomaticFunctionCallingConfig(disable=True),
        # 思考の要約も受け取る(利用者が、どう考えて SQL にしたかを確かめながら進めるため)
        thinking_config=types.ThinkingConfig(thinking_level=types.ThinkingLevel(level), include_thoughts=True),
    )
    contents: list[Any] = [types.Content(role=m.role, parts=[types.Part(text=m.text)]) for m in body.messages]
    text = ""
    try:
        for round_ in range(MAX_TOOL_ROUNDS + 1):
            if round_ == MAX_TOOL_ROUNDS:
                # 道具を使う回数の上限。ここまでで分かったことで答えさせる(道具を渡さずに、もう 1 回)
                contents.append(types.Content(role="user", parts=[types.Part(text="道具を使う回数の上限に達しました。ここまでで分かったことで、決まりの JSON で答えてください。")]))
                config = config.model_copy(update={"tools": None, "automatic_function_calling": None})
            parts: list[Any] = []
            calls: list[Any] = []
            text = ""
            stream = await ai._get_client().aio.models.generate_content_stream(model=model, contents=contents, config=config)
            async for chunk in stream:
                for cand in chunk.candidates or []:
                    for p in (cand.content.parts if cand.content else None) or []:
                        if p.thought:
                            if p.text:
                                yield {"type": "thought", "text": p.text}
                            # 思考の署名(thought_signature)は、次の呼び出しに戻す必要がある
                            if p.thought_signature:
                                parts.append(p)
                            continue
                        parts.append(p)
                        if p.function_call:
                            calls.append(p.function_call)
                        elif p.text:
                            text += p.text
            if not calls:
                break
            # 返答(思考の署名を含む)をそのまま履歴に戻してから、道具の結果を足す
            contents.append(types.Content(role="model", parts=parts))
            results = []
            for call in calls:
                out, step = await asyncio.to_thread(run_tool, client, call.name or "", dict(call.args or {}), body)
                yield {"type": "step", "text": step}
                results.append(types.Part.from_function_response(name=call.name or "", response=out))
            contents.append(types.Content(role="user", parts=results))
    except Exception as e:  # noqa: BLE001 — Vertex AI のエラーは種類が多い。原因をそのまま見せる
        yield {"type": "error", "error": f"{model} の呼び出しに失敗しました: {type(e).__name__}: {e}"[:2000]}
        return
    text = text.strip()
    if not text:
        yield {"type": "error", "error": "返答が空でした(安全フィルターなどで止まった可能性があります)"}
        return
    yield {"type": "final", **parse_reply(text), "context_chars": len(context)}


def _ndjson(events: AsyncIterator[dict[str, Any]]) -> StreamingResponse:
    """起きた順に、1 行 1 つの JSON(NDJSON)で返す。考えている間も、15 秒ごとに合図を送る
    (中継(Jupyter のプロキシ)が、黙っている接続を切らないように。画面も「止まっていない」と分かる)"""

    async def gen() -> AsyncIterator[bytes]:
        it = events.__aiter__()
        pending = asyncio.ensure_future(it.__anext__())
        while True:
            done, _ = await asyncio.wait({pending}, timeout=15)
            if not done:
                yield _line({"type": "ping"})
                continue
            try:
                ev = pending.result()
            except StopAsyncIteration:
                return
            yield _line(ev)
            pending = asyncio.ensure_future(it.__anext__())

    # X-Accel-Buffering: 途中の中継が、まとめて送らないように
    return StreamingResponse(gen(), media_type="application/x-ndjson", headers={"Cache-Control": "no-cache", "X-Accel-Buffering": "no"})


@router.post("/ai/chat")
async def chat(body: ChatIn) -> StreamingResponse:
    """会話の 1 回分(思考・調べたこと・返事を、起きた順に)"""
    return _ndjson(chat_events(body))


@router.post("/ai/context")
async def context_preview(body: ChatIn) -> dict[str, Any]:
    """Gemini に渡す前提を、そのまま返す(画面の「前提を見る・コピー」。DEVTOOLS.md の 3.1 ③)"""
    try:
        return {"ok": True, "context": (await _context(body))[0]}
    except Exception as e:  # noqa: BLE001
        return _error(e)


# ---- 出力ごとの作業: 版を残す・答え合わせ(DEVTOOLS.md の 3.1 の ④〜⑦)
# 版は server/devwork/outputs/<出力>/v<番号>.sql と .json(作った日時・メモ・答え合わせの結果)。
# 採用の印は、同じフォルダの index.json。DB に置かないのは、import-pack で DB を置き換えても消えないように。

OUTPUT_NAMES_ALL = tuple(o.name for o in bq.OUTPUTS)
# 全件の答え合わせで読む行の上限(これを超える出力は、集計の誤りを疑う。メモリを守るため)
CHECK_MAX_ROWS = 2_000_000


def _work_dir(name: str) -> Path:
    if name not in OUTPUT_NAMES_ALL:
        raise HTTPException(404, f"知らない出力です: {name}")
    return DEVWORK / "outputs" / name


def _read_json(p: Path, fallback: Any) -> Any:
    try:
        return json.loads(p.read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return fallback


def _write_json(p: Path, v: Any) -> None:
    p.parent.mkdir(parents=True, exist_ok=True)
    p.write_text(json.dumps(v, ensure_ascii=False, indent=2, default=str), encoding="utf-8")


def versions(name: str) -> list[dict[str, Any]]:
    d = _work_dir(name)
    out = []
    for p in sorted(d.glob("v*.json"), key=lambda x: int(re.sub(r"\D", "", x.stem) or 0)):
        meta = _read_json(p, {})
        meta["v"] = int(re.sub(r"\D", "", p.stem) or 0)
        out.append(meta)
    return out


def _now_iso() -> str:
    return datetime.now().astimezone().isoformat(timespec="seconds")


@router.get("/work")
async def work() -> dict[str, Any]:
    """6 つの出力の版の一覧・採用の版と、出す先の設定"""
    outs = {}
    for o in bq.OUTPUTS:
        outs[o.name] = {"label": o.label, "env": o.env, "versions": versions(o.name), "adopted": _read_json(_work_dir(o.name) / "index.json", {}).get("adopted")}
    return {"ok": True, "outputs": outs, "target": _read_json(DEVWORK / "target.json", {})}


class TargetIn(BaseModel):
    # 出す先のデータセット(<プロジェクト>.<データセット>)と、表の名前の頭(ccp_ なら ccp_spu)
    dataset: str = Field("", max_length=300)
    prefix: str = Field("ccp_", max_length=60)


@router.put("/work/target")
async def put_target(body: TargetIn) -> dict[str, Any]:
    _write_json(DEVWORK / "target.json", body.model_dump())
    return {"ok": True}


class VersionIn(BaseModel):
    sql: str = Field(min_length=1, max_length=200_000)
    note: str = Field("", max_length=2000)


@router.post("/work/{name}")
async def save_version(name: str, body: VersionIn) -> dict[str, Any]:
    d = _work_dir(name)
    d.mkdir(parents=True, exist_ok=True)
    have = versions(name)
    # 直前の版と同じ SQL なら、版を増やさない(保存を何度押しても散らからないように)
    if have:
        last = have[-1]["v"]
        try:
            if (d / f"v{last}.sql").read_text(encoding="utf-8").strip() == body.sql.strip():
                return {"ok": True, "v": last, "same": True}
        except OSError:
            pass
    v = (have[-1]["v"] if have else 0) + 1
    (d / f"v{v}.sql").write_text(body.sql, encoding="utf-8")
    _write_json(d / f"v{v}.json", {"created": _now_iso(), "note": body.note, "check": None})
    return {"ok": True, "v": v, "same": False}


@router.get("/work/{name}/{v}")
async def get_version(name: str, v: int) -> dict[str, Any]:
    d = _work_dir(name)
    try:
        sql = (d / f"v{v}.sql").read_text(encoding="utf-8")
    except OSError:
        raise HTTPException(404, f"{name} の v{v} はありません")
    return {"ok": True, "v": v, "sql": sql, **_read_json(d / f"v{v}.json", {})}


class AdoptIn(BaseModel):
    v: int | None = None


@router.put("/work/{name}/adopt")
async def adopt(name: str, body: AdoptIn) -> dict[str, Any]:
    _write_json(_work_dir(name) / "index.json", {"adopted": body.v})
    return {"ok": True}


class CheckIn(BaseModel):
    output: str
    sql: str = Field(min_length=1, max_length=200_000)
    # 答え合わせの結果を、この版に残す(なければ残さない)
    version: int | None = None
    # 結びつき(SKU → SPU など)も確かめる。相手の出力の「採用の版」を流して、キーだけ読む
    with_context: bool = True
    # 読む量を見てから流す: false なら dry-run だけ(読む量・料金の概算を返す)
    run: bool = False


# 結びつきを確かめる相手(bq_validate.validate の context)
CONTEXT_OF = {"sku": ("spu",), "sku_stock": ("sku",), "weekly": ("spu",)}
CONTEXT_KEYS = {"spu": ("market", "product_code"), "sku": ("market", "sku_code")}


def _adopted_sql(name: str) -> str | None:
    v = _read_json(_work_dir(name) / "index.json", {}).get("adopted")
    if not v:
        return None
    try:
        return (_work_dir(name) / f"v{v}.sql").read_text(encoding="utf-8")
    except OSError:
        return None


def _strip_sql(sql: str) -> str:
    """副問い合わせに包めるように、末尾の ; を外す"""
    return re.sub(r";\s*$", "", sql.strip())


def check_output(c: Any, body: CheckIn) -> dict[str, Any]:
    from google.cloud import bigquery

    from . import bq_validate as bv

    if body.output not in OUTPUT_NAMES_ALL:
        return {"ok": False, "error": f"知らない出力です: {body.output}"}
    o = next(x for x in bq.OUTPUTS if x.name == body.output)
    dry = _dry_run(c, body.sql)
    if (why := only_select(dry.statement_type)) is not None:
        return {"ok": False, "error": why}
    cols = [f.name for f in (dry.schema or [])]
    missing, missing_optional = bq.compare_columns(o, cols)
    # 結びつきの相手(採用の版があるものだけ)
    ctx_sqls = {}
    if body.with_context:
        for other in CONTEXT_OF.get(body.output, ()):
            s = _adopted_sql(other)
            if s:
                keys = ", ".join(CONTEXT_KEYS[other])
                ctx_sqls[other] = f"SELECT DISTINCT {keys} FROM (\n{_strip_sql(s)}\n)"
    est = dry.total_bytes_processed or 0
    for s in ctx_sqls.values():
        try:
            est += _dry_run(c, s).total_bytes_processed or 0
        except Exception:  # noqa: BLE001 — 相手の版が流せないなら、結びつきは確かめない(下で伝える)
            pass
    base = {"ok": True, "output": body.output, "columns": cols, "missing": missing, "missing_optional": missing_optional,
            "bytes": est, "cost_usd": cost_usd(max(est, 10 * 1024**2)), "price_per_tib": price_per_tib(),
            "weight": weight(est, None), "context": sorted(ctx_sqls), "max_bytes": max_bytes()}
    if not body.run:
        return {**base, "ran": False}
    if missing:
        return {**base, "ran": False, "error": f"必須の列がありません: {', '.join(missing)}(列の名前は、仕様のとおりに AS で付け替えてください)"}
    cfg = bigquery.QueryJobConfig(maximum_bytes_billed=max_bytes())
    job = c.query(body.sql, job_config=cfg)
    it = job.result()
    if (it.total_rows or 0) > CHECK_MAX_ROWS:
        return {**base, "ran": False, "error": f"{it.total_rows:,} 行あります。答え合わせは {CHECK_MAX_ROWS:,} 行までです(集計の粒度を見直してください)"}
    rows = [{k.lower(): v for k, v in dict(r).items()} for r in it]
    context: dict[str, list[dict[str, Any]]] = {}
    notes: list[str] = []
    for other, s in ctx_sqls.items():
        try:
            context[other] = [{k.lower(): v for k, v in dict(r).items()} for r in c.query(s, job_config=cfg).result()]
        except Exception as e:  # noqa: BLE001
            notes.append(f"{other} の採用の版を流せなかったので、結びつきは確かめていません({_error(e)['error'][:200]})")
    for other in CONTEXT_OF.get(body.output, ()):
        if body.with_context and other not in ctx_sqls:
            notes.append(f"{other} に採用の版がないので、{other} との結びつきは確かめていません")
    issues = bv.validate(body.output, rows, context)
    out_issues = [{"level": i.level, "message": i.message, "count": i.count, "examples": [str(x) for x in i.examples]} for i in issues]
    for col in missing_optional:
        out_issues.append({"level": bv.WARN, "message": f"あると使う列 {col} がありません", "count": 0, "examples": []})
    ng = sum(1 for i in out_issues if i["level"] == bv.NG)
    warn = len(out_issues) - ng
    result = {"checked_at": _now_iso(), "rows": len(rows), "ng": ng, "warn": warn}
    if body.version:
        meta_p = _work_dir(body.output) / f"v{body.version}.json"
        if meta_p.is_file():
            meta = _read_json(meta_p, {})
            meta["check"] = result
            _write_json(meta_p, meta)
    return {**base, "ran": True, "rows": len(rows), "issues": out_issues, "ng": ng, "warn": warn, "notes": notes,
            "checked_at": result["checked_at"], "job": job_info(c, body.sql, job)}


@router.post("/check")
async def check(body: CheckIn) -> dict[str, Any]:
    """答え合わせ。run=false で読む量の見積もりだけ、run=true で全件を流して bq_validate にかける(CCP の取り込みと同じ確かめ)"""
    try:
        return await asyncio.to_thread(check_output, bq.client(), body)
    except Exception as e:  # noqa: BLE001
        return _error(e)


# ---- AI との会話(セッション)を残す: server/devwork/chats/<id>.json
# ブラウザに残すだけだと、別のブラウザ・PC から続きを話せないので、Workbench のフォルダに置く(Git には入らない)

CHAT_ID = re.compile(r"^[A-Za-z0-9_-]{6,64}$")


def _chat_path(cid: str) -> Path:
    if not CHAT_ID.match(cid):
        raise HTTPException(400, "会話の ID の形が違います")
    return DEVWORK / "chats" / f"{cid}.json"


def chat_title(messages: list[dict[str, Any]]) -> str:
    """名前がないときは、最初の発言の頭から付ける"""
    first = next((str(m.get("text") or "") for m in messages if m.get("role") == "user"), "")
    t = re.sub(r"\s+", " ", first).strip()
    return (t[:40] + "…") if len(t) > 40 else (t or "新しい会話")


@router.get("/chats")
async def chats() -> dict[str, Any]:
    """会話の一覧(新しく話した順)。中身は返さない"""
    out = []
    for p in (DEVWORK / "chats").glob("*.json"):
        v = _read_json(p, None)
        if not isinstance(v, dict):
            continue
        msgs = v.get("messages") or []
        out.append({"id": p.stem, "title": v.get("title") or chat_title(msgs), "created": v.get("created"), "updated": v.get("updated"),
                    "count": len(msgs), "output": v.get("output") or ""})
    out.sort(key=lambda x: x.get("updated") or "", reverse=True)
    return {"ok": True, "chats": out}


@router.get("/chats/{cid}")
async def get_chat(cid: str) -> dict[str, Any]:
    v = _read_json(_chat_path(cid), None)
    if not isinstance(v, dict):
        raise HTTPException(404, "その会話はありません")
    return {"ok": True, "id": cid, **v}


class ChatSaveIn(BaseModel):
    title: str = Field("", max_length=200)
    output: str = Field("", max_length=40)
    # 画面の形のまま(思考・調べたこと・SQL も含めて)残す
    messages: list[dict[str, Any]] = Field(default_factory=list, max_length=500)


@router.put("/chats/{cid}")
async def save_chat(cid: str, body: ChatSaveIn) -> dict[str, Any]:
    p = _chat_path(cid)
    old = _read_json(p, {}) if p.is_file() else {}
    now = _now_iso()
    _write_json(p, {"title": body.title.strip() or old.get("title") or chat_title(body.messages), "output": body.output,
                    "created": old.get("created") or now, "updated": now, "messages": body.messages})
    return {"ok": True}


@router.delete("/chats/{cid}")
async def delete_chat(cid: str) -> dict[str, Any]:
    p = _chat_path(cid)
    if p.is_file():
        p.unlink()
    return {"ok": True}


# ---- ポリシーの文書(DEVTOOLS.md の 3.2): GCS の文書を、CCP の決まりの形(POLICY_FORMAT.md)に書き直す
# GCS は読むだけ。確かめ・取り込みは、既存の /api/policies/preview・/api/policies/import を画面から呼ぶ(取り込みと同じ読み分け)

GCS = "https://storage.googleapis.com/storage/v1"
# 1 つのファイルの上限(Gemini にそのまま渡す PDF も、これに収める)
MAX_FILE = 20 * 1024**2
# 文字にした中身を Gemini に渡す上限(1 ファイル)
MAX_TEXT = 300_000
TEXT_EXT = (".md", ".txt", ".csv", ".tsv", ".json")


def _gcs_get(path: str, params: dict[str, Any] | None = None, raw: bool = False) -> Any:
    from .assets import session

    r = session().get(f"{GCS}{path}", params=params, timeout=60)
    if r.status_code >= 400:
        raise RuntimeError(f"GCS {r.status_code}: {r.text[:300]}")
    return r.content if raw else r.json()


def _obj_path(bucket: str, name: str) -> str:
    from urllib.parse import quote

    return f"/b/{quote(bucket, safe='')}/o/{quote(name, safe='')}"


@router.get("/gcs/list")
async def gcs_list(bucket: str, prefix: str = "") -> dict[str, Any]:
    """バケットの、フォルダ 1 段分の一覧(フォルダと、ファイル)"""
    def go() -> dict[str, Any]:
        b = bucket.strip().removeprefix("gs://").strip("/")
        out = {"ok": True, "bucket": b, "prefix": prefix, "folders": [], "files": []}
        token = None
        while True:
            params = {"prefix": prefix, "delimiter": "/", "maxResults": 1000, **({"pageToken": token} if token else {})}
            v = _gcs_get(f"/b/{b}/o", params)
            out["folders"] += v.get("prefixes", [])
            out["files"] += [{"name": o["name"], "size": int(o.get("size", 0)), "updated": o.get("updated"), "type": o.get("contentType", "")}
                             for o in v.get("items", []) if not o["name"].endswith("/")]
            token = v.get("nextPageToken")
            if not token or len(out["files"]) >= 3000:
                break
        return out

    try:
        return await asyncio.to_thread(go)
    except Exception as e:  # noqa: BLE001
        return _error(e)


def decode_text(data: bytes) -> str:
    """文字の文書。UTF-8 で読めなければ Shift_JIS(Excel で書き出した CSV に多い)"""
    for enc in ("utf-8-sig", "cp932"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def xlsx_text(data: bytes) -> str:
    """Excel を、シートごとの CSV の形の文字にする(見出しはシート名)。
    結合したセルは左上の値で埋める(結合した見出しの下の行が空にならないように)。数式は、計算済みの値を読む"""
    import csv
    import io

    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), data_only=True)
    parts = []
    for ws in wb.worksheets:
        if ws.sheet_state != "visible":
            continue
        fill: dict[tuple[int, int], Any] = {}
        for rng in ws.merged_cells.ranges:
            v = ws.cell(rng.min_row, rng.min_col).value
            for r in range(rng.min_row, rng.max_row + 1):
                for c in range(rng.min_col, rng.max_col + 1):
                    fill[(r, c)] = v
        rows = []
        for r_idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
            vals = [fill.get((r_idx, c_idx), v) for c_idx, v in enumerate(row, start=1)]
            cells = ["" if v is None else (v.isoformat() if hasattr(v, "isoformat") else str(v)).strip() for v in vals]
            while cells and not cells[-1]:
                cells.pop()
            if any(cells):
                rows.append(cells)
        if not rows:
            continue
        # 空の列(どの行にも値がない列)を落とす
        width = max(len(r) for r in rows)
        keep = [c for c in range(width) if any(c < len(r) and r[c] for r in rows)]
        buf = io.StringIO()
        w = csv.writer(buf, lineterminator="\n")
        for r in rows:
            w.writerow([r[c] if c < len(r) else "" for c in keep])
        parts.append(f"### シート: {ws.title}\n\n```csv\n{buf.getvalue().rstrip()}\n```")
    return "\n\n".join(parts) or "(値のあるシートがありません)"


def kind_of(name: str) -> str:
    n = name.lower()
    if n.endswith((".xlsx", ".xlsm")):
        return "xlsx"
    if n.endswith(".pdf"):
        return "pdf"
    if n.endswith(TEXT_EXT):
        return "text"
    return "other"


def read_doc(bucket: str, name: str) -> dict[str, Any]:
    """GCS の文書 1 つを読み、Gemini に渡せる形にする。PDF は中身をそのまま(Gemini が読める)"""
    k = kind_of(name)
    if k == "other":
        return {"name": name, "kind": k, "error": "この形式は読めません(Excel・PDF・CSV・テキストだけ。Word は PDF に書き出してください)"}
    meta = _gcs_get(_obj_path(bucket, name))
    size = int(meta.get("size", 0))
    if size > MAX_FILE:
        return {"name": name, "kind": k, "error": f"大きすぎます({size / 1024**2:.1f} MB。{MAX_FILE // 1024**2} MB まで)"}
    data = _gcs_get(_obj_path(bucket, name), {"alt": "media"}, raw=True)
    if k == "pdf":
        return {"name": name, "kind": k, "bytes": data, "size": size}
    text = xlsx_text(data) if k == "xlsx" else decode_text(data)
    cut = len(text) > MAX_TEXT
    return {"name": name, "kind": k, "text": text[:MAX_TEXT], "size": size, "chars": len(text), "truncated": cut}


class DocsIn(BaseModel):
    bucket: str = Field(min_length=1, max_length=300)
    names: list[str] = Field(min_length=1, max_length=20)


@router.post("/gcs/extract")
async def gcs_extract(body: DocsIn) -> dict[str, Any]:
    """選んだ文書を文字にして返す(画面で中身を確かめるため。PDF は文字にしない)"""
    def go() -> dict[str, Any]:
        out = []
        for n in body.names:
            try:
                d = read_doc(body.bucket.strip().removeprefix("gs://").strip("/"), n)
            except Exception as e:  # noqa: BLE001
                d = {"name": n, "kind": kind_of(n), "error": _error(e)["error"]}
            d.pop("bytes", None)
            out.append(d)
        return {"ok": True, "docs": out}

    try:
        return await asyncio.to_thread(go)
    except Exception as e:  # noqa: BLE001
        return _error(e)


POLICY_ROLE = """あなたは、マーケティングポリシー・文言のルールの文書を、Campaign Contents Planner(CCP)に取り込める形に書き直す専門家です。
渡された文書(Excel をシートごとに CSV の形にしたもの・PDF・テキスト)を読み、下の「取り込むファイルの形」(POLICY_FORMAT.md)のとおりに書き直してください。

返答は、次の形の JSON だけにしてください(```json で囲まない):
{"reply": "利用者への返事(日本語。Markdown を使ってよい。何をどのファイルに分けたか・決めきれなかったこと)",
 "files": [{"kind": "rules または sales", "filename": "取り込むときのファイル名(rules は .md、sales は .csv。元の文書の名前から)", "content": "ファイルの中身の全文"}],
 "questions": ["利用者に確かめたいこと(なければ空の配列)"]}

決まり:
- ルール系(表現・選定・口調などの決まり)は Markdown(rules)、販売系の方針(週 × 市場 × 商品やカテゴリの方針)は CSV(sales)。1 つの文書に両方あれば、2 つのファイルに分ける
- 文書にないことを足さない。読み取れない・解釈が分かれるところは questions で聞く
- 決まりごとに、元の文書のどこから取ったか(シート名・ページ・見出し)を、POLICY_FORMAT.md で使える項目(出どころ・メモなど)に書く。項目がなければ reply に書く
- 個人の名前は書かない(役割にする)
- 「確かめの結果」(読めなかった決まり・行と理由)が渡されたら、それを直した全文を返す
"""


class ConvertIn(BaseModel):
    bucket: str = Field(min_length=1, max_length=300)
    names: list[str] = Field(min_length=1, max_length=20)
    # 利用者からの頼み(どの市場の決まりか・どのシートを使うか など)。続けて話すときは、これまでの会話も
    messages: list[ChatTurn] = Field(default_factory=list, max_length=100)
    model: str | None = Field(None, max_length=200)


async def convert_events(body: ConvertIn) -> AsyncIterator[dict[str, Any]]:
    """文書を書き直す 1 回分を、起きた順に返す(会話と同じ形: start / step / thought / final / error)"""
    from google.genai import types

    from . import ai

    model = body.model or await default_model()
    level = (os.getenv("DEVTOOLS_THINKING") or "high").upper()
    yield {"type": "start", "model": model, "thinking": level.lower()}
    bucket = body.bucket.strip().removeprefix("gs://").strip("/")
    parts: list[Any] = []
    for n in body.names:
        try:
            d = await asyncio.to_thread(read_doc, bucket, n)
        except Exception as e:  # noqa: BLE001
            yield {"type": "error", "error": f"{n} を読めません: {_error(e)['error']}"}
            return
        if d.get("error"):
            yield {"type": "error", "error": f"{n}: {d['error']}"}
            return
        if d["kind"] == "pdf":
            parts.append(types.Part(text=f"## 文書: {n}(PDF)"))
            parts.append(types.Part.from_bytes(data=d["bytes"], mime_type="application/pdf"))
            yield {"type": "step", "text": f"{n} を読んだ(PDF、{d['size'] / 1024:.0f} KB。そのまま渡す)"}
        else:
            parts.append(types.Part(text=f"## 文書: {n}\n\n{d['text']}"))
            yield {"type": "step", "text": f"{n} を文字にした({d['chars']:,} 文字{'。長いので途中まで' if d.get('truncated') else ''})"}
    system = POLICY_ROLE + "\n\n## 取り込むファイルの形(POLICY_FORMAT.md)\n\n" + _doc("POLICY_FORMAT.md")
    turns = body.messages or [ChatTurn(role="user", text="これらの文書を、取り込める形に書き直してください。")]
    # 文書は最初の発言に添える(続けて話すときも、毎回渡す。前の返事の中身は会話に入っている)
    contents = [types.Content(role="user", parts=[*parts, types.Part(text=turns[0].text)])]
    contents += [types.Content(role=t.role, parts=[types.Part(text=t.text)]) for t in turns[1:]]
    config = types.GenerateContentConfig(
        system_instruction=system,
        response_mime_type="application/json",
        thinking_config=types.ThinkingConfig(thinking_level=types.ThinkingLevel(level), include_thoughts=True),
    )
    text = ""
    try:
        stream = await ai._get_client().aio.models.generate_content_stream(model=model, contents=contents, config=config)
        async for chunk in stream:
            for cand in chunk.candidates or []:
                for p in (cand.content.parts if cand.content else None) or []:
                    if p.thought:
                        if p.text:
                            yield {"type": "thought", "text": p.text}
                    elif p.text:
                        text += p.text
    except Exception as e:  # noqa: BLE001
        yield {"type": "error", "error": f"{model} の呼び出しに失敗しました: {type(e).__name__}: {e}"[:2000]}
        return
    v = parse_files(text)
    if not v["files"] and not v["reply"]:
        yield {"type": "error", "error": "返答が空でした(安全フィルターなどで止まった可能性があります)"}
        return
    # 書き直したファイルは、作業フォルダにも残す(画面を閉じても、あとで見られるように)
    for f in v["files"]:
        try:
            p = DEVWORK / "policies" / re.sub(r"[\\/:*?\"<>|]", "_", f["filename"])
            p.parent.mkdir(parents=True, exist_ok=True)
            p.write_text(f["content"], encoding="utf-8")
        except OSError:
            pass
    yield {"type": "final", **v}


def parse_files(text: str) -> dict[str, Any]:
    """書き直しの返答(JSON)を読む。崩れていたら、本文を返事として扱う"""
    for c in (text, text[text.find("{"): text.rfind("}") + 1] if "{" in text else ""):
        try:
            v = json.loads(c)
        except ValueError:
            continue
        if isinstance(v, dict):
            files = []
            for f in v.get("files") or []:
                if not isinstance(f, dict) or not str(f.get("content") or "").strip():
                    continue
                kind = "sales" if str(f.get("kind")).lower() == "sales" else "rules"
                name = str(f.get("filename") or "").strip() or f"policy.{'csv' if kind == 'sales' else 'md'}"
                files.append({"kind": kind, "filename": name, "content": str(f["content"])})
            q = v.get("questions")
            return {"reply": str(v.get("reply") or ""), "files": files, "questions": [str(x) for x in q if x] if isinstance(q, list) else []}
    return {"reply": text, "files": [], "questions": []}


@router.post("/policy/convert")
async def policy_convert(body: ConvertIn) -> StreamingResponse:
    """文書を書き直す(NDJSON。会話の /ai/chat と同じ形で、15 秒ごとに合図を送る)"""
    return _ndjson(convert_events(body))


@router.get("/policy/drafts")
async def policy_drafts() -> dict[str, Any]:
    """前に書き直したファイル(server/devwork/policies/)の一覧と中身"""
    out = []
    for p in sorted((DEVWORK / "policies").glob("*"), key=lambda x: x.stat().st_mtime, reverse=True):
        if p.is_file():
            out.append({"filename": p.name, "kind": "sales" if p.suffix.lower() == ".csv" else "rules", "content": p.read_text(encoding="utf-8"),
                        "updated": datetime.fromtimestamp(p.stat().st_mtime).astimezone().isoformat(timespec="seconds")})
    return {"ok": True, "drafts": out}


# ---- 作業ファイル(server/devwork/files/): 引き継ぎ・SQL・コーディングルールなど。AI も道具で読み書きする
# 「いつも渡す」の印を付けたファイルは、会話の前提に毎回入れる(コーディングルールを守らせる・引き継ぎから続ける)

FILE_NAME = re.compile(r"^[^\\/:*?\"<>|\x00-\x1f]{1,120}\.(md|sql|txt|csv)$", re.I)
# いつも渡すファイルの、前提に入れる量の上限(全部で)
PINNED_MAX = 200_000


def _files_dir() -> Path:
    return DEVWORK / "files"


def _file_path(name: str) -> Path:
    name = name.strip()
    if not FILE_NAME.match(name) or name.startswith("."):
        raise HTTPException(400, f"ファイル名は .md / .sql / .txt / .csv で、/ や \\ を含めない名前にしてください: {name}")
    return _files_dir() / name


def pinned_names() -> list[str]:
    v = _read_json(_files_dir() / ".pinned.json", [])
    return [x for x in v if isinstance(x, str)] if isinstance(v, list) else []


def list_work_files() -> list[dict[str, Any]]:
    pins = set(pinned_names())
    out = []
    for p in _files_dir().glob("*"):
        if p.is_file() and not p.name.startswith(".") and FILE_NAME.match(p.name):
            st = p.stat()
            out.append({"name": p.name, "size": st.st_size, "updated": datetime.fromtimestamp(st.st_mtime).astimezone().isoformat(timespec="seconds"),
                        "pinned": p.name in pins})
    out.sort(key=lambda x: x["updated"], reverse=True)
    return out


def read_work_file(name: str) -> str:
    p = _file_path(name)
    if not p.is_file():
        raise HTTPException(404, f"{name} はありません")
    return p.read_text(encoding="utf-8")


def write_work_file(name: str, content: str) -> bool:
    """保存する。もとからあったか(上書きしたか)を返す"""
    p = _file_path(name)
    existed = p.is_file()
    p.parent.mkdir(parents=True, exist_ok=True)
    p.write_text(content, encoding="utf-8")
    return existed


def pinned_context() -> str:
    """いつも渡すファイルを、前提の 1 節にする(量の上限を超えたら、後ろを切って、切ったと書く)"""
    parts = []
    used = 0
    for n in pinned_names():
        try:
            text = read_work_file(n)
        except Exception:  # noqa: BLE001 — 消されたファイルは飛ばす
            continue
        room = PINNED_MAX - used
        if room <= 0:
            parts.append(f"### {n}\n(量の上限を超えたので、渡していません)")
            continue
        cut = text[:room]
        used += len(cut)
        parts.append(f"### {n}\n\n{cut}" + ("\n\n(量の上限で、ここまで)" if len(text) > room else ""))
    if not parts:
        return ""
    return "## いつも渡すファイル(守ること・引き継ぎ。利用者が印を付けたもの。必ず読んで従う)\n\n" + "\n\n".join(parts)


@router.get("/files")
async def files() -> dict[str, Any]:
    return {"ok": True, "files": list_work_files()}


@router.get("/files/{name}")
async def get_file(name: str) -> dict[str, Any]:
    return {"ok": True, "name": name, "content": read_work_file(name), "pinned": name in pinned_names()}


class FileIn(BaseModel):
    content: str = Field("", max_length=2_000_000)


@router.put("/files/{name}")
async def put_file(name: str, body: FileIn) -> dict[str, Any]:
    existed = write_work_file(name, body.content)
    return {"ok": True, "overwritten": existed}


@router.delete("/files/{name}")
async def delete_file(name: str) -> dict[str, Any]:
    p = _file_path(name)
    if p.is_file():
        p.unlink()
    _write_json(_files_dir() / ".pinned.json", [x for x in pinned_names() if x != name])
    return {"ok": True}


class PinIn(BaseModel):
    pinned: bool


@router.put("/files/{name}/pin")
async def pin_file(name: str, body: PinIn) -> dict[str, Any]:
    _file_path(name)
    pins = [x for x in pinned_names() if x != name]
    if body.pinned:
        pins.append(name)
    _write_json(_files_dir() / ".pinned.json", pins)
    return {"ok": True}


class XlsxIn(BaseModel):
    # PC から上げた Excel(base64)。コーディングルールが Excel のときに、シートごとの表の形の Markdown にして置く
    data: str = Field(min_length=1, max_length=30_000_000)


@router.post("/files/{name}/from-xlsx")
async def file_from_xlsx(name: str, body: XlsxIn) -> dict[str, Any]:
    try:
        text = await asyncio.to_thread(xlsx_text, base64.b64decode(body.data))
    except Exception as e:  # noqa: BLE001
        return _error(e)
    existed = write_work_file(name, text)
    return {"ok": True, "overwritten": existed, "chars": len(text)}
