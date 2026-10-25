"""開発用の画面の API で、どの窓口も使う小道具: 有効か・作業フォルダ・BigQuery の値の形と料金の目安・JSON のファイル・文書を文字にする"""
from __future__ import annotations

import base64
import json
import os
from datetime import date, datetime, time as dtime
from decimal import Decimal
from pathlib import Path
from typing import Any

from fastapi import HTTPException


DEVWORK = Path(os.getenv("DEVTOOLS_DIR") or Path(__file__).resolve().parents[2] / "devwork")
# 画面に返す行の上限(多すぎるとブラウザが重くなる。全件の確かめは、別の段で行う)
MAX_ROWS = 1000
PREVIEW_ROWS = 20


def enabled() -> bool:
    return os.getenv("CCP_DEVTOOLS", "").strip().lower() in ("1", "true", "yes", "on")


def blocks(path: str) -> bool:
    """無効のとき、このパスを隠すか(画面の /dev/ と、API の /api/dev/)。main.py のミドルウェアから呼ぶ"""
    return not enabled() and (path == "/dev" or path.startswith("/dev/") or path.startswith("/api/dev/"))


def max_bytes() -> int:
    try:
        gb = float(os.getenv("DEVTOOLS_MAX_GB") or 10)
    except ValueError:
        gb = 10
    return int(gb * 1024**3)


async def require_enabled() -> None:
    if not enabled():
        raise HTTPException(404)


def plain(v: Any) -> Any:
    """BigQuery の値を、JSON にできる形にする(日時は ISO の文字、数値の Decimal は文字のまま桁を落とさない、バイト列は base64)"""
    if v is None or isinstance(v, (bool, int, float, str)):
        return v
    if isinstance(v, (datetime, date, dtime)):
        return v.isoformat()
    if isinstance(v, Decimal):
        return str(v)
    if isinstance(v, bytes):
        return base64.b64encode(v).decode()
    if isinstance(v, dict):
        return {k: plain(x) for k, x in v.items()}
    if isinstance(v, (list, tuple)):
        return [plain(x) for x in v]
    return str(v)


def flatten_fields(fields: Any, prefix: str = "") -> list[dict[str, Any]]:
    """表の列(SchemaField)を、入れ子(RECORD)を a.b の形に開いて並べる。Gemini に渡す列の一覧にもそのまま使う"""
    out: list[dict[str, Any]] = []
    for f in fields or []:
        name = f"{prefix}{f.name}"
        out.append({"name": name, "type": f.field_type, "mode": f.mode or "NULLABLE", "description": f.description or ""})
        if getattr(f, "fields", None):
            out.extend(flatten_fields(f.fields, f"{name}."))
    return out


def schema_columns(fields: Any) -> list[dict[str, str]]:
    return [{"name": f.name, "type": f.field_type} for f in (fields or [])]


def only_select(statement_type: str | None) -> str | None:
    """流してよい文か。だめなら、理由を返す(画面にそのまま出す)"""
    if statement_type == "SELECT":
        return None
    if not statement_type:
        return "文の種類が分かりませんでした(複数の文・スクリプトは流せません)"
    return f"SELECT ではないので流しません(文の種類: {statement_type})。表を作る SQL は、BigQuery のコンソールで実行してください"


def error_result(e: Exception) -> dict[str, Any]:
    # BigQuery のエラーは、本文(message)に行と列の位置が入っている。Gemini に返して直させるので、そのまま渡す
    msg = getattr(e, "message", None) or str(e)
    errors = getattr(e, "errors", None) or []
    detail = "; ".join(x.get("message", "") for x in errors if isinstance(x, dict) and x.get("message")) if errors else ""
    return {"ok": False, "error": f"{type(e).__name__}: {detail or msg}"[:4000]}


def dry_run(client: Any, sql: str) -> Any:
    from google.cloud import bigquery

    return client.query(sql, job_config=bigquery.QueryJobConfig(dry_run=True, use_query_cache=False))


# ---- SQL を試す

def price_per_tib() -> float:
    """オンデマンドの料金(1 TiB あたりのドル)。場所や契約で違うので、設定(DEVTOOLS_PRICE_PER_TIB)で変えられる"""
    try:
        return float(os.getenv("DEVTOOLS_PRICE_PER_TIB") or 6.25)
    except ValueError:
        return 6.25


def cost_usd(n: int | None) -> float | None:
    return None if n is None else round(n / 1024**4 * price_per_tib(), 4)


def weight(processed: int | None, slot_ms: int | None) -> str:
    """重さの目安(読む量とスロット時間から。画面に一言で出す)"""
    gb = (processed or 0) / 1024**3
    slot_s = (slot_ms or 0) / 1000
    if gb >= 100 or slot_s >= 3600:
        return "重い"
    if gb >= 1 or slot_s >= 60:
        return "ふつう"
    return "軽い"


DOCS = Path(__file__).resolve().parents[3] / "docs"


def doc(name: str) -> str:
    try:
        return (DOCS / name).read_text(encoding="utf-8")
    except OSError:
        return ""


def read_json(p: Path, fallback: Any) -> Any:
    try:
        return json.loads(p.read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return fallback


def write_json(p: Path, v: Any) -> None:
    p.parent.mkdir(parents=True, exist_ok=True)
    p.write_text(json.dumps(v, ensure_ascii=False, indent=2, default=str), encoding="utf-8")


def now_iso() -> str:
    return datetime.now().astimezone().isoformat(timespec="seconds")


def decode_text(data: bytes) -> str:
    """文字の文書。UTF-8 で読めなければ Shift_JIS(Excel で書き出した CSV に多い)"""
    for enc in ("utf-8-sig", "cp932"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def xlsx_text(data: bytes) -> str:
    """Excel を、シートごとの CSV の形の文字にする(見出しはシート名)。
    結合したセルは左上の値で埋める(結合した見出しの下の行が空にならないように)。数式は、計算済みの値を読む"""
    import csv
    import io

    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), data_only=True)
    parts = []
    for ws in wb.worksheets:
        if ws.sheet_state != "visible":
            continue
        fill: dict[tuple[int, int], Any] = {}
        for rng in ws.merged_cells.ranges:
            v = ws.cell(rng.min_row, rng.min_col).value
            for r in range(rng.min_row, rng.max_row + 1):
                for c in range(rng.min_col, rng.max_col + 1):
                    fill[(r, c)] = v
        rows = []
        for r_idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
            vals = [fill.get((r_idx, c_idx), v) for c_idx, v in enumerate(row, start=1)]
            cells = ["" if v is None else (v.isoformat() if hasattr(v, "isoformat") else str(v)).strip() for v in vals]
            while cells and not cells[-1]:
                cells.pop()
            if any(cells):
                rows.append(cells)
        if not rows:
            continue
        # 空の列(どの行にも値がない列)を落とす
        width = max(len(r) for r in rows)
        keep = [c for c in range(width) if any(c < len(r) and r[c] for r in rows)]
        buf = io.StringIO()
        w = csv.writer(buf, lineterminator="\n")
        for r in rows:
            w.writerow([r[c] if c < len(r) else "" for c in keep])
        parts.append(f"### シート: {ws.title}\n\n```csv\n{buf.getvalue().rstrip()}\n```")
    return "\n\n".join(parts) or "(値のあるシートがありません)"
