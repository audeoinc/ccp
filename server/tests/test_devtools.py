"""devtools.py(開発用の画面の API)の単体テスト。BigQuery には接続しない。

  cd server && python -m unittest discover tests
"""
from __future__ import annotations

import sys
import unittest
from datetime import date, datetime, timezone
from decimal import Decimal
from pathlib import Path
from types import SimpleNamespace
from unittest import mock

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from ccp import devtools  # noqa: E402


class EnabledTest(unittest.TestCase):
    def test_hidden_unless_enabled(self) -> None:
        with mock.patch.dict("os.environ", {"CCP_DEVTOOLS": ""}):
            self.assertTrue(devtools.blocks("/dev/"))
            self.assertTrue(devtools.blocks("/dev"))
            self.assertTrue(devtools.blocks("/api/dev/status"))
            # CCP の画面・API は隠さない(/devices のような、たまたま dev で始まる名前も)
            self.assertFalse(devtools.blocks("/"))
            self.assertFalse(devtools.blocks("/api/data/products"))
            self.assertFalse(devtools.blocks("/devices"))
        with mock.patch.dict("os.environ", {"CCP_DEVTOOLS": "1"}):
            self.assertFalse(devtools.blocks("/dev/"))
            self.assertFalse(devtools.blocks("/api/dev/status"))


class OnlySelectTest(unittest.TestCase):
    def test_select_only(self) -> None:
        self.assertIsNone(devtools.only_select("SELECT"))
        for st in ("CREATE_TABLE_AS_SELECT", "INSERT", "DELETE", "SCRIPT", "DROP_TABLE"):
            self.assertIsNotNone(devtools.only_select(st), st)
        self.assertIsNotNone(devtools.only_select(None))


class PlainTest(unittest.TestCase):
    def test_values(self) -> None:
        row = {
            "ts": datetime(2026, 10, 22, 9, 0, tzinfo=timezone.utc),
            "d": date(2026, 10, 19),
            "n": Decimal("12345678901234567890.12"),
            "b": b"\x00\x01",
            "arr": [Decimal("1.5"), None],
            "rec": {"x": date(2026, 1, 1)},
        }
        out = devtools.plain(row)
        self.assertEqual(out["ts"], "2026-10-22T09:00:00+00:00")
        self.assertEqual(out["d"], "2026-10-19")
        # 桁を落とさない(float にすると丸まる)
        self.assertEqual(out["n"], "12345678901234567890.12")
        self.assertEqual(out["b"], "AAE=")
        self.assertEqual(out["arr"], ["1.5", None])
        self.assertEqual(out["rec"], {"x": "2026-01-01"})


class FlattenTest(unittest.TestCase):
    def test_nested(self) -> None:
        f = lambda name, t, sub=None, mode="NULLABLE": SimpleNamespace(name=name, field_type=t, mode=mode, description="", fields=sub or [])  # noqa: E731
        cols = devtools.flatten_fields([f("id", "STRING", mode="REQUIRED"), f("item", "RECORD", [f("color", "STRING"), f("size", "STRING")], "REPEATED")])
        self.assertEqual([c["name"] for c in cols], ["id", "item", "item.color", "item.size"])
        self.assertEqual(cols[1]["mode"], "REPEATED")


class ChatTest(unittest.TestCase):
    def test_chapters(self) -> None:
        md = "# 題\n\n## 1. 全体\nA\n\n---\n\n## 2. SPU\nB\n\n## 3. SKU\nC\n"
        self.assertEqual(devtools.chapters(md, (1, 3)), "## 1. 全体\nA\n\n## 3. SKU\nC")

    def test_parse_reply(self) -> None:
        self.assertEqual(devtools.parse_reply('{"reply": "直しました", "sql": " SELECT 1 ", "questions": ["q", ""]}'), {"reply": "直しました", "sqls": [{"title": "", "name": "", "sql": "SELECT 1"}], "questions": ["q"]})
        # JSON が崩れていても、本文を返事として見せる
        self.assertEqual(devtools.parse_reply("ただの文")["reply"], "ただの文")

    def test_private_columns(self) -> None:
        for name in ("customer_name", "item_nm", "email", "tel_no", "review_text", "氏名"):
            self.assertTrue(devtools.PRIVATE.search(name), name)
        for name in ("sku_cd", "qty", "sales_date", "price"):
            self.assertFalse(devtools.PRIVATE.search(name), name)

    def test_context_has_spec_and_sql(self) -> None:
        body = devtools.ChatIn(output="sku", sql="SELECT 1", last="BadRequest: x", notes="まとめ")
        c = devtools.build_context(body, None)
        self.assertIn("## 3. SKU", c)
        self.assertIn("## 1. 全体のルール", c)
        self.assertIn("SELECT 1", c)
        self.assertIn("まとめ", c)


class ToolTest(unittest.TestCase):
    def client(self) -> SimpleNamespace:
        tables = [SimpleNamespace(project="p", dataset_id="a", table_id=n, table_type="TABLE") for n in ("m_ctlg_item", "CTLG_price", "t_sales")]
        return SimpleNamespace(project="p", list_tables=lambda ds: tables if ds == "p.a" else [],
                               query=lambda sql, job_config=None: SimpleNamespace(statement_type="DELETE", total_bytes_processed=0, schema=[]))

    def test_list_tables_by_part_of_name(self) -> None:
        out, step = devtools.run_tool(self.client(), "list_tables", {"dataset": "a", "name_contains": "ctlg"}, devtools.ChatIn())
        self.assertEqual([t["table"] for t in out["tables"]], ["p.a.m_ctlg_item", "p.a.CTLG_price"])
        self.assertIn("2 件", step)

    def test_run_select_only(self) -> None:
        out, _ = devtools.run_tool(self.client(), "run_select", {"sql": "DELETE FROM x WHERE true"}, devtools.ChatIn())
        self.assertIn("SELECT ではない", out["error"])
        out, _ = devtools.run_tool(self.client(), "run_select", {"sql": "SELECT 1"}, devtools.ChatIn(allow_run=False))
        self.assertIn("止めて", out["error"])

    def test_tools_declared(self) -> None:
        names = [f.name for f in devtools.tool_declarations()[0].function_declarations]
        self.assertEqual(names, ["list_datasets", "list_tables", "describe_table", "describe_view_sources", "dry_run", "profile_column", "list_files", "read_file", "save_file", "run_select"])

    def test_view_sources(self) -> None:
        refs = [SimpleNamespace(project="p", dataset_id="v", table_id="vw"), SimpleNamespace(project="p", dataset_id="raw", table_id="item"),
                SimpleNamespace(project="p", dataset_id="raw", table_id="item"), SimpleNamespace(project="p", dataset_id="raw", table_id="price")]
        c = SimpleNamespace(query=lambda sql, job_config=None: SimpleNamespace(referenced_tables=refs))
        # ビュー自身と重なりは除く
        self.assertEqual(devtools.view_sources(c, "p.v.vw"), ["p.raw.item", "p.raw.price"])

    def test_referenced_views(self) -> None:
        # ジョブの記録には元の表だけ。SQL に書いたビューは、名前から見つける
        job = SimpleNamespace(referenced_tables=[SimpleNamespace(project="p", dataset_id="raw", table_id="item")])
        kinds = {"p.mart.v_item": "VIEW", "p.raw.item": "TABLE"}

        def get_table(name: str) -> SimpleNamespace:
            if name not in kinds:
                raise KeyError(name)
            pr, ds, t = name.split(".")
            return SimpleNamespace(project=pr, dataset_id=ds, table_id=t, table_type=kinds[name])

        c = SimpleNamespace(get_table=get_table)
        r = devtools.referenced(c, "SELECT a.x FROM `p.mart.v_item` a JOIN p.raw.item b ON a.id = b.id", job)
        self.assertEqual(r, {"tables": ["p.raw.item"], "views": ["p.mart.v_item"]})

    def test_weight_and_cost(self) -> None:
        self.assertEqual(devtools.weight(50 * 1024**2, 2000), "軽い")
        self.assertEqual(devtools.weight(5 * 1024**3, None), "ふつう")
        self.assertEqual(devtools.weight(200 * 1024**3, None), "重い")
        self.assertAlmostEqual(devtools.cost_usd(1024**4), 6.25)

    def test_parse_fenced(self) -> None:
        r = devtools.parse_reply('調べました。\n```json\n{"reply": "a", "sql": "SELECT 1"}\n```')
        self.assertEqual((r["reply"], r["sqls"][0]["sql"]), ("a", "SELECT 1"))

    def test_parse_many_sqls(self) -> None:
        r = devtools.parse_reply('{"reply": "2 つ", "sqls": [{"title": "値の確かめ", "sql": "SELECT DISTINCT a FROM t"}, {"title": "本番", "sql": "SELECT a FROM t"}, {"title": "空", "sql": ""}]}')
        self.assertEqual([x["title"] for x in r["sqls"]], ["値の確かめ", "本番"])

    def test_pro_of(self) -> None:
        self.assertEqual(devtools.pro_of("gemini-3.8-flash"), "gemini-3.8-pro")
        self.assertEqual(devtools.pro_of("gemini-3.5-flash-lite"), "gemini-3.5-pro")
        self.assertIsNone(devtools.pro_of("gemini-3.8-pro"))

    def test_pick_models(self) -> None:
        names = ["publishers/google/models/gemini-3.5-flash", "publishers/google/models/gemini-3.5-pro", "publishers/google/models/gemini-3.8-pro-preview",
                 "publishers/google/models/gemini-3.8-pro", "publishers/google/models/gemini-3.5-flash-image", "publishers/google/models/text-embedding-005",
                 "publishers/google/models/gemini-3.5-flash-lite"]
        self.assertEqual(devtools.pick_models(names), ["gemini-3.8-pro", "gemini-3.8-pro-preview", "gemini-3.5-pro", "gemini-3.5-flash"])


class WorkTest(unittest.TestCase):
    def setUp(self) -> None:
        import tempfile

        self.tmp = tempfile.TemporaryDirectory()
        self.patch = mock.patch.object(devtools, "DEVWORK", Path(self.tmp.name))
        self.patch.start()

    def tearDown(self) -> None:
        self.patch.stop()
        self.tmp.cleanup()

    def run_async(self, coro):  # noqa: ANN001, ANN202
        import asyncio

        return asyncio.run(coro)

    def test_versions_and_adopt(self) -> None:
        a = self.run_async(devtools.save_version("spu", devtools.VersionIn(sql="SELECT 1")))
        same = self.run_async(devtools.save_version("spu", devtools.VersionIn(sql="SELECT 1 ")))
        b = self.run_async(devtools.save_version("spu", devtools.VersionIn(sql="SELECT 2", note="直した")))
        # 同じ SQL では版を増やさない
        self.assertEqual((a["v"], same["v"], same["same"], b["v"]), (1, 1, True, 2))
        self.run_async(devtools.adopt("spu", devtools.AdoptIn(v=2)))
        w = self.run_async(devtools.work())
        self.assertEqual([x["v"] for x in w["outputs"]["spu"]["versions"]], [1, 2])
        self.assertEqual(w["outputs"]["spu"]["adopted"], 2)
        self.assertEqual(self.run_async(devtools.get_version("spu", 2))["sql"], "SELECT 2")
        self.assertEqual(devtools._adopted_sql("spu"), "SELECT 2")

    def test_chats(self) -> None:
        msgs = [{"role": "user", "text": "ctlg を含む表から  CCP で使える列を探して"}, {"role": "model", "text": "見ました", "thoughts": "…"}]
        self.run_async(devtools.save_chat("abc123", devtools.ChatSaveIn(messages=msgs, output="spu")))
        self.run_async(devtools.save_chat("def456", devtools.ChatSaveIn(title="名前つき", messages=[])))
        lst = self.run_async(devtools.chats())["chats"]
        self.assertEqual({c["id"]: c["title"] for c in lst}, {"abc123": "ctlg を含む表から CCP で使える列を探して", "def456": "名前つき"})
        got = self.run_async(devtools.get_chat("abc123"))
        self.assertEqual((got["output"], got["messages"][1]["thoughts"]), ("spu", "…"))
        # 名前を付け直さずに残しても、前の名前のまま
        self.run_async(devtools.save_chat("def456", devtools.ChatSaveIn(messages=msgs)))
        self.assertEqual(self.run_async(devtools.get_chat("def456"))["title"], "名前つき")
        self.run_async(devtools.delete_chat("abc123"))
        self.assertEqual([c["id"] for c in self.run_async(devtools.chats())["chats"]], ["def456"])
        with self.assertRaises(Exception):
            devtools._chat_path("../x")

    def test_work_files_and_pinned_context(self) -> None:
        self.run_async(devtools.put_file("コーディングルール.md", devtools.FileIn(content="# SQL の決まり\n- 予約語は大文字")))
        self.run_async(devtools.put_file("spu.sql", devtools.FileIn(content="SELECT 1")))
        self.run_async(devtools.pin_file("コーディングルール.md", devtools.PinIn(pinned=True)))
        names = {f["name"]: f["pinned"] for f in self.run_async(devtools.files())["files"]}
        self.assertEqual(names, {"コーディングルール.md": True, "spu.sql": False})
        # いつも渡すファイルは、前提の頭のほう(仕様より前)に入る
        c = devtools.build_context(devtools.ChatIn(), None)
        self.assertIn("### コーディングルール.md", c)
        self.assertLess(c.index("## いつも渡すファイル"), c.index("## CCP が読む出力の仕様"))
        self.assertNotIn("SELECT 1", c)
        # AI の道具でも読み書きできる(上書きかどうかも返す)
        out, step = devtools.run_tool(None, "save_file", {"name": "引き継ぎ.md", "content": "## 決まったこと"}, devtools.ChatIn())
        self.assertEqual((out["overwritten"], devtools.read_work_file("引き継ぎ.md")), (False, "## 決まったこと"))
        out, _ = devtools.run_tool(None, "save_file", {"name": "../外.md", "content": "x"}, devtools.ChatIn())
        self.assertIn("error", out)
        self.run_async(devtools.delete_file("コーディングルール.md"))
        self.assertEqual(devtools.pinned_names(), [])

    def test_history(self) -> None:
        devtools.record_run("画面", "SELECT 1", rows=1, processed=100, billed=10485760, seconds=0.5)
        devtools.record_run("AI", "SELECT bad", error="BadRequest: x")
        runs = devtools.read_history()
        # 新しい順。失敗も残る
        self.assertEqual([(r["source"], r["ok"]) for r in runs], [("AI", False), ("画面", True)])
        self.assertEqual([r["sql"] for r in devtools.read_history(q="bad")], ["SELECT bad"])

    def test_profile_column(self) -> None:
        f = SimpleNamespace(name="country_cd", field_type="STRING", mode="NULLABLE", description="", fields=[])
        t = SimpleNamespace(project="p", dataset_id="d", table_id="t", schema=[f])
        big = {"n": 2 * 1024**3}

        def query(sql, job_config=None):  # noqa: ANN001, ANN202
            if getattr(job_config, "dry_run", False):
                return SimpleNamespace(total_bytes_processed=big["n"])
            rows = [SimpleNamespace(items=lambda: {"total": 4, "nulls": 1, "distinct_approx": 2, "min_value": "JPN", "max_value": "USA"}.items())] if "COUNT(*) AS total" in sql                 else [SimpleNamespace(items=lambda: {"value": "JPN", "n": 2}.items()), SimpleNamespace(items=lambda: {"value": "USA", "n": 1}.items())]
            return SimpleNamespace(result=lambda: rows)

        c = SimpleNamespace(get_table=lambda ref: t, query=query)
        # 1 GB を超えると、確かめてから
        self.assertTrue(devtools.profile_column(c, "p.d.t", "country_cd", confirm=False).get("needs_confirm"))
        r = devtools.profile_column(c, "p.d.t", "country_cd", confirm=True)
        self.assertEqual((r["total"], r["nulls"], r["null_rate"], [v["value"] for v in r["top"]]), (4, 1, 0.25, ["JPN", "USA"]))
        self.assertEqual(devtools.read_history()[0]["source"], "列を調べる")
        self.assertFalse(devtools.profile_column(c, "p.d.t", "nope", confirm=True)["ok"])

    def test_check_runs_validate_with_context(self) -> None:
        self.run_async(devtools.save_version("spu", devtools.VersionIn(sql="SELECT spu")))
        self.run_async(devtools.adopt("spu", devtools.AdoptIn(v=1)))
        f = lambda n: SimpleNamespace(name=n)  # noqa: E731
        schema = [f(n) for n in ("market", "sku_code", "product_code", "as_of")]
        sku_rows = [{"market": "JP", "sku_code": "S1", "product_code": "P1", "as_of": datetime.now(timezone.utc)},
                    {"market": "JP", "sku_code": "S2", "product_code": "P9", "as_of": datetime.now(timezone.utc)}]
        spu_rows = [{"market": "JP", "product_code": "P1"}]
        sqls: list[str] = []

        def query(sql: str, job_config=None):  # noqa: ANN001, ANN202
            sqls.append(sql)
            if getattr(job_config, "dry_run", False):
                return SimpleNamespace(statement_type="SELECT", total_bytes_processed=1000, schema=schema)
            rows = spu_rows if "SELECT spu" in sql else sku_rows
            job = SimpleNamespace(result=lambda: _Rows(rows), job_id="j", location="x", total_bytes_processed=1000, total_bytes_billed=10485760,
                                  slot_millis=10, started=None, ended=None, cache_hit=False, referenced_tables=[])
            return job

        c = SimpleNamespace(query=query, get_table=lambda n: (_ for _ in ()).throw(KeyError(n)))
        r = devtools.check_output(c, devtools.CheckIn(output="sku", sql="SELECT sku", run=True))
        self.assertTrue(r["ran"], r)
        self.assertEqual(r["context"], ["spu"])
        # P9 は SPU の採用の版にない → NG
        self.assertTrue(any("SPU の出力にありません" in i["message"] for i in r["issues"] if i["level"] == "NG"), r["issues"])
        self.assertGreaterEqual(r["ng"], 1)
        # 結びつきの相手は、採用の版をキーだけにして流す
        self.assertTrue(any("SELECT DISTINCT market, product_code FROM (" in q for q in sqls))


class CancelTest(unittest.TestCase):
    def test_cancel(self) -> None:
        import asyncio

        cancelled = []
        devtools._running["r1"] = ("job_1", "asia-northeast1")
        with mock.patch.object(devtools.bq, "client", return_value=SimpleNamespace(cancel_job=lambda j, location=None: cancelled.append((j, location)))):
            r = asyncio.run(devtools.cancel(devtools.CancelIn(run_id="r1")))
            done = asyncio.run(devtools.cancel(devtools.CancelIn(run_id="nope")))
        self.assertEqual((r["cancelled"], cancelled), (True, [("job_1", "asia-northeast1")]))
        self.assertFalse(done["cancelled"])
        devtools._running.clear()


class TaskTest(unittest.TestCase):
    def test_tasks_run_in_background(self) -> None:
        import asyncio

        async def main() -> list:
            async def slow() -> dict:
                await asyncio.sleep(0.05)
                return {"ok": True, "rows": 3}

            async def events():  # noqa: ANN202
                yield {"type": "thought", "text": "a"}
                raise RuntimeError("落ちた")

            t1 = devtools.start_call(slow())
            t2 = devtools.start_events(events())
            first = await devtools.task_status(t1)
            await asyncio.sleep(0.1)
            done = await devtools.task_status(t1)
            err = await devtools.task_status(t2)
            more = await devtools.task_status(t2, after=1)
            return [first, done, err, more, await devtools.task_status("nope")]

        first, done, err, more, nope = asyncio.run(main())
        # 受付はすぐ返り、あとで結果が届く。途中で落ちても、エラーとして届く
        self.assertFalse(first["done"])
        self.assertEqual((done["done"], done["events"][0]["result"]["rows"]), (True, 3))
        self.assertEqual([e["type"] for e in err["events"]], ["thought", "error"])
        self.assertEqual([e["type"] for e in more["events"]], ["error"])
        self.assertFalse(nope["ok"])

    def test_task_call_rejects_unknown(self) -> None:
        import asyncio

        r = asyncio.run(devtools.task_call(devtools.CallIn(path="files", body={})))
        self.assertFalse(r["ok"])


class PolicyDocTest(unittest.TestCase):
    def test_xlsx_text(self) -> None:
        import io

        from openpyxl import Workbook

        wb = Workbook()
        ws = wb.active
        ws.title = "表現の決まり"
        ws.append(["区分", "決まり", None, "メモ"])
        ws.append(["価格", "税込みで書く", None, ""])
        ws.append([None, "割引率は整数", None, None])
        ws.merge_cells("A2:A3")
        ws.append([None, None, None, None])
        hidden = wb.create_sheet("下書き")
        hidden.append(["見せない"])
        hidden.sheet_state = "hidden"
        buf = io.BytesIO()
        wb.save(buf)
        t = devtools.xlsx_text(buf.getvalue())
        self.assertIn("### シート: 表現の決まり", t)
        # 結合したセルは左上の値で埋める。空の列(C)は落とす。隠したシートは読まない
        self.assertIn("区分,決まり,メモ", t)
        self.assertIn("価格,割引率は整数,", t)
        self.assertNotIn("見せない", t)

    def test_decode_and_kind(self) -> None:
        self.assertEqual(devtools.decode_text("市場,方針".encode("cp932")), "市場,方針")
        self.assertEqual(devtools.decode_text("﻿a,b".encode("utf-8")), "a,b")
        self.assertEqual([devtools.kind_of(n) for n in ("a.XLSX", "b.pdf", "c.csv", "d.docx")], ["xlsx", "pdf", "text", "other"])

    def test_parse_files(self) -> None:
        v = devtools.parse_files('{"reply": "2 つに分けました", "files": [{"kind": "rules", "filename": "表現.md", "content": "# ルール集"}, {"kind": "SALES", "filename": "", "content": "a,b"}, {"kind": "rules", "content": ""}], "questions": []}')
        self.assertEqual([(f["kind"], f["filename"]) for f in v["files"]], [("rules", "表現.md"), ("sales", "policy.csv")])
        self.assertEqual(devtools.parse_files("文だけ")["reply"], "文だけ")

    def test_gcs_list(self) -> None:
        import asyncio

        pages = [{"prefixes": ["policy/2026/"], "items": [{"name": "policy/", "size": "0"}, {"name": "policy/a.xlsx", "size": "120", "updated": "t", "contentType": "x"}], "nextPageToken": "n"},
                 {"items": [{"name": "policy/b.pdf", "size": "9"}]}]
        with mock.patch.object(devtools, "_gcs_get", side_effect=pages):
            r = asyncio.run(devtools.gcs_list("gs://bkt/", "policy/"))
        self.assertEqual((r["bucket"], r["folders"], [f["name"] for f in r["files"]]), ("bkt", ["policy/2026/"], ["policy/a.xlsx", "policy/b.pdf"]))


class _Rows(list):
    def __init__(self, rows):  # noqa: ANN001
        super().__init__(rows)
        self.total_rows = len(rows)


if __name__ == "__main__":
    unittest.main()
